from __future__ import annotations

import hashlib
from pathlib import Path

import pytest

from rnaseq_helper.recognition.scanner import AssetScanner
from rnaseq_helper.preparation.service import PreparationService
from rnaseq_helper.readers.geo_reader import GeoReadCancelled
from rnaseq_helper.readers.table_reader import TableReadCancelled, TableReader
from rnaseq_helper.ui.workers.import_worker import ImportTaskWorker


@pytest.mark.parametrize("filename", ["reads.fastq.gz", "reads.fq.gz"])
def test_scan_recognizes_compressed_fastq_as_raw_sequencing(
    tmp_path: Path, filename: str
) -> None:
    source = tmp_path / filename
    source.write_bytes(b"placeholder compressed sequencing data")

    result = AssetScanner().scan((source,))

    assert len(result.assets) == 1
    asset = result.assets[0]
    assert asset.kind == "RAW_SEQUENCING"
    assert asset.role == "GUIDANCE_ONLY"
    assert asset.source_paths == (source.resolve(),)


def test_scan_keeps_every_selected_file_in_the_manifest(tmp_path: Path) -> None:
    sources = (tmp_path / "first.pdf", tmp_path / "second.pdf")
    for source in sources:
        source.write_bytes(b"%PDF-1.4 auxiliary document")

    result = AssetScanner().scan(sources)

    discovered = {path for asset in result.assets for path in asset.source_paths}
    assert discovered == {source.resolve() for source in sources}
    assert len(result.assets) == 2


def test_mixed_delivery_folder_keeps_pdf_counts_and_fastq_separate(tmp_path: Path) -> None:
    package = tmp_path / "delivery"
    package.mkdir()
    (package / "a.pdf").write_bytes(b"%PDF-1.4 delivery note")
    (package / "b_counts.tsv").write_text(
        "raw counts\tcontrol_1\tcontrol_2\nfeature_a\t12\t14\nfeature_b\t21\t19\n",
        encoding="utf-8",
    )
    (package / "sample.fastq.gz").write_bytes(b"raw sequencing placeholder")

    result = AssetScanner().scan((package,))

    by_name = {asset.source_paths[0].name: asset for asset in result.assets}
    assert set(by_name) == {"a.pdf", "b_counts.tsv", "sample.fastq.gz"}
    assert by_name["a.pdf"].kind == "AUXILIARY"
    assert by_name["b_counts.tsv"].kind == "TABLE"
    assert by_name["b_counts.tsv"].measurement_hint.value == "RAW_COUNT"
    assert by_name["sample.fastq.gz"].kind == "RAW_SEQUENCING"
    assert "不会作为表达矩阵读取" in " ".join(item.basis for item in by_name["a.pdf"].evidence)


def test_tenx_bundle_does_not_swallow_other_delivery_files(tmp_path: Path) -> None:
    package = tmp_path / "tenx_delivery"
    package.mkdir()
    (package / "matrix.mtx").write_text(
        "%%MatrixMarket matrix coordinate integer general\n"
        "2 2 2\n1 1 3\n2 2 5\n",
        encoding="utf-8",
    )
    (package / "barcodes.tsv").write_text("cell1\ncell2\n", encoding="utf-8")
    (package / "features.tsv").write_text(
        "gene1\tG1\tGene Expression\ngene2\tG2\tGene Expression\n",
        encoding="utf-8",
    )
    (package / "counts.tsv").write_text(
        "raw counts\tcontrol_1\tcontrol_2\nfeature_a\t12\t14\nfeature_b\t21\t19\n",
        encoding="utf-8",
    )
    (package / "notes.pdf").write_bytes(b"%PDF-1.4 delivery note")
    (package / "sample.fastq.gz").write_bytes(b"raw sequencing placeholder")

    result = AssetScanner().scan((package,))

    tenx = next(asset for asset in result.assets if asset.kind == "TENX")
    by_name = {
        asset.source_paths[0].name: asset
        for asset in result.assets
        if len(asset.source_paths) == 1
    }
    assert set(tenx.source_paths) == {
        (package / name).resolve()
        for name in ("matrix.mtx", "barcodes.tsv", "features.tsv")
    }
    assert by_name["counts.tsv"].kind == "TABLE"
    assert by_name["notes.pdf"].kind == "AUXILIARY"
    assert by_name["sample.fastq.gz"].kind == "RAW_SEQUENCING"


def test_incomplete_tenx_bundle_only_groups_recognized_candidates(tmp_path: Path) -> None:
    package = tmp_path / "partial_tenx_delivery"
    package.mkdir()
    matrix = package / "matrix.mtx"
    matrix.write_text("%%MatrixMarket matrix coordinate integer general\n2 2 0\n", encoding="utf-8")
    (package / "counts.tsv").write_text(
        "feature\tsample_a\tsample_b\ngene_a\t3\t5\n",
        encoding="utf-8",
    )
    (package / "notes.pdf").write_bytes(b"%PDF-1.4 delivery note")
    (package / "sample.fastq.gz").write_bytes(b"raw sequencing placeholder")

    result = AssetScanner().scan((package,))

    tenx = next(asset for asset in result.assets if asset.kind == "TENX")
    by_name = {
        asset.source_paths[0].name: asset
        for asset in result.assets
        if len(asset.source_paths) == 1
    }
    assert tenx.source_paths == (matrix.resolve(),)
    assert tenx.role == "UNREADABLE"
    assert any(issue.code == "TENX_INCOMPLETE" for issue in tenx.issues)
    assert by_name["counts.tsv"].kind == "TABLE"
    assert by_name["notes.pdf"].kind == "AUXILIARY"
    assert by_name["sample.fastq.gz"].kind == "RAW_SEQUENCING"


def test_workbook_recommends_raw_counts_over_tpm_without_confirming_unit() -> None:
    assert AssetScanner.recommended_sheet(("TPM", "raw_counts")) == "raw_counts"
    assert AssetScanner.recommended_sheet(("normalized_counts", "TPM")) == "TPM"
    assert AssetScanner.recommended_sheet(("counts_A", "counts_B", "TPM")) is None
    assert AssetScanner.recommended_sheet(("Sheet A", "Sheet B")) is None


def test_pdf_is_classified_as_ignorable_auxiliary_file(tmp_path: Path) -> None:
    source = tmp_path / "protocol.pdf"
    source.write_bytes(b"%PDF-1.4")

    result = AssetScanner().scan((source,))

    assert len(result.assets) == 1
    assert result.assets[0].kind == "AUXILIARY"
    assert result.assets[0].role == "IGNORE_CANDIDATE"


def test_scan_does_not_hash_every_input_in_full(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = tmp_path / "reads.fastq"
    source.write_bytes(b"@read\nACGT\n+\n!!!!\n")
    hash_calls = 0
    original_sha256 = hashlib.sha256

    def tracked_sha256(*args: object, **kwargs: object) -> object:
        nonlocal hash_calls
        hash_calls += 1
        return original_sha256(*args, **kwargs)

    monkeypatch.setattr(hashlib, "sha256", tracked_sha256)

    AssetScanner().scan((source,))

    assert hash_calls == 0, "扫描不应为输入文件计算完整内容 SHA256"


def test_scan_reports_observable_progress(tmp_path: Path) -> None:
    sources = tuple(tmp_path / f"reads_{index}.fastq" for index in range(3))
    for source in sources:
        source.write_bytes(b"@read\nACGT\n+\n!!!!\n")
    progress_events = []

    result = AssetScanner().scan(sources, progress=progress_events.append)

    assert len(result.assets) == len(sources)
    assert progress_events, "扫描应至少报告一次进度"
    assert all(event.phase for event in progress_events)
    assert all(event.total == len(sources) for event in progress_events)
    completed = [event.completed for event in progress_events]
    assert completed == sorted(completed)
    assert completed[-1] == len(sources)
    assert all(event.path for event in progress_events)


def test_directory_discovery_progress_does_not_repeat_after_file_count_stalls(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    root = tmp_path / "many_entries"
    root.mkdir()
    files = []
    for index in range(25):
        source = root / f"note_{index:02d}.pdf"
        source.write_bytes(b"%PDF-1.4")
        files.append(source)
    directories = []
    for index in range(50):
        directory = root / f"empty_{index:02d}"
        directory.mkdir()
        directories.append(directory)

    original_rglob = Path.rglob

    def deterministic_rglob(path: Path, pattern: str):
        if path == root:
            return iter([*files, *directories])
        return original_rglob(path, pattern)

    monkeypatch.setattr(Path, "rglob", deterministic_rglob)
    progress_events = []

    result = AssetScanner().scan((root,), progress=progress_events.append)

    discovery_events = [event for event in progress_events if event.phase == "discover"]
    assert len(result.assets) == len(files)
    assert len(discovery_events) == 3


def test_full_table_read_reports_progress_and_stops_cooperatively(tmp_path: Path) -> None:
    source = tmp_path / "large.tsv"
    source.write_text(
        "feature\tsample_a\tsample_b\n"
        + "".join(f"gene_{index}\t{index}\t{index + 1}\n" for index in range(1500)),
        encoding="utf-8",
    )
    cancel_requested = False
    progress_events: list[int] = []

    def report_progress(rows_read: int) -> None:
        nonlocal cancel_requested
        progress_events.append(rows_read)
        cancel_requested = True

    with pytest.raises(TableReadCancelled):
        TableReader().read(
            source,
            progress=report_progress,
            cancelled=lambda: cancel_requested,
        )

    assert progress_events == [1000]


def test_table_reader_marks_bounded_preview_and_complete_read(tmp_path: Path) -> None:
    source = tmp_path / "matrix.tsv"
    source.write_text(
        "feature\tsample_a\tsample_b\n"
        + "".join(f"gene_{index}\t{index}\t{index + 1}\n" for index in range(150)),
        encoding="utf-8",
    )

    preview = TableReader().read(source, max_rows=40, max_columns=80)
    complete = TableReader().read(source)

    assert preview.row_count == 40
    assert preview.is_complete is False
    assert complete.row_count == 151
    assert complete.is_complete is True


def test_preparation_preview_defers_full_read_until_explicit_completion(tmp_path: Path) -> None:
    source = tmp_path / "counts.tsv"
    source.write_text(
        "feature\tsample_a\tsample_b\n"
        + "".join(f"gene_{index}\t{index}\t{index + 1}\n" for index in range(150)),
        encoding="utf-8",
    )
    asset = AssetScanner().scan((source,)).assets[0]
    service = PreparationService()

    preview_session = service.start(asset, tmp_path / "project", preview_only=True)
    complete_session = service.complete_read(preview_session)

    assert preview_session.raw_table.row_count == 40
    assert preview_session.raw_table.is_complete is False
    assert preview_session.preview_shape == (40, 3)
    assert complete_session.raw_table.row_count == 151
    assert complete_session.raw_table.is_complete is True
    assert complete_session.recipe == preview_session.recipe
    assert complete_session.preview_shape == preview_session.preview_shape


def test_validation_blocks_columns_that_were_outside_the_preview(tmp_path: Path) -> None:
    source = tmp_path / "wide.tsv"
    headers = ["feature", *[f"sample_{index}" for index in range(170)]]
    rows = [headers, ["gene_a", *["12" for _ in range(170)]], ["gene_b", *["15" for _ in range(170)]]]
    source.write_text("\n".join("\t".join(row) for row in rows) + "\n", encoding="utf-8")
    asset = AssetScanner().scan((source,)).assets[0]
    service = PreparationService()
    preview_session = service.start(asset, tmp_path / "project", preview_only=True)
    assert preview_session.preview_shape == (3, 80)
    confirmed = service.apply(preview_session, {
        "measurement_type": "RAW_COUNT",
        "value_scale": "LINEAR",
    })
    confirmed_complete = service.complete_read(confirmed)

    summary = service.validate(confirmed_complete)

    assert any(issue.code == "PREVIEW_LIMIT_EXCEEDED" for issue in summary.issues)


def test_preview_covers_a_realistic_wide_geo_sample_axis(tmp_path: Path) -> None:
    source = tmp_path / "wide_series_matrix.txt"
    headers = ["ID_REF", *[f"GSM{index:06d}" for index in range(1, 145)]]
    rows = [headers, ["probe_a", *[str(index) for index in range(144)]]]
    source.write_text(
        "!Sample_geo_accession = " + "\t".join(headers[1:]) + "\n"
        "!series_matrix_table_begin\n"
        + "\n".join("\t".join(row) for row in rows)
        + "\n!series_matrix_table_end\n",
        encoding="utf-8",
    )

    scanner = AssetScanner()
    asset = scanner.scan((source,)).assets[0]
    session = PreparationService().start(asset, tmp_path / "project", preview_only=True)
    preview = PreparationService().preview(session, window=(0, 5))

    assert session.preview_shape == (2, 145)
    assert session.raw_table.is_complete is True
    assert preview.column_count_hint == 145
    assert len(preview.rows[0]) == 145
    assert preview.rows[0][-1] == "GSM000144"


def test_geo_preview_is_bounded_and_followup_full_read_is_cancellable(tmp_path: Path) -> None:
    source = tmp_path / "series_matrix.txt"
    source.write_text(
        "!Sample_geo_accession = GSM1\n"
        "!series_matrix_table_begin\n"
        "ID_REF\tGSM1\n"
        + "".join(f"probe_{index}\t{index}\n" for index in range(1500))
        + "!series_matrix_table_end\n",
        encoding="utf-8",
    )
    scanner = AssetScanner()
    asset = scanner.scan((source,)).assets[0]
    assert asset.kind == "GEO_SERIES_MATRIX"
    assert asset.shape_hint == (40, 2)

    service = PreparationService()
    session = service.start(asset, tmp_path / "project", preview_only=True)
    assert session.raw_table.row_count == 40
    assert session.raw_table.is_complete is False

    cancel_requested = False
    phases: list[str] = []

    def report(rows_read: int) -> None:
        nonlocal cancel_requested
        phases.append(str(rows_read))
        if rows_read >= 1000:
            cancel_requested = True

    with pytest.raises(GeoReadCancelled):
        service.complete_read(session, progress=report, cancelled=lambda: cancel_requested)

    assert phases == ["1000"]


def test_scan_returns_cancelled_result_when_cancellation_is_requested(
    tmp_path: Path,
) -> None:
    sources = tuple(tmp_path / f"reads_{index}.fastq" for index in range(3))
    for source in sources:
        source.write_bytes(b"@read\nACGT\n+\n!!!!\n")

    result = AssetScanner().scan(sources, cancelled=lambda: True)

    assert result.cancelled is True


def test_cancelled_scan_keeps_completed_candidates_in_worker_result(tmp_path: Path) -> None:
    sources = tuple(tmp_path / f"reads_{index}.fastq" for index in range(3))
    for source in sources:
        source.write_text("@read\nACGT\n+\n!!!!\n", encoding="utf-8")
    worker = ImportTaskWorker(
        lambda progress, cancelled: AssetScanner().scan(
            sources, progress=progress, cancelled=cancelled
        )
    )
    succeeded = []
    cancelled_events = []
    worker.progress.connect(lambda _event: worker.request_cancel())
    worker.succeeded.connect(succeeded.append)
    worker.cancelled.connect(lambda: cancelled_events.append(True))

    worker.run()

    assert cancelled_events == []
    assert len(succeeded) == 1
    assert succeeded[0].cancelled is True
    assert len(succeeded[0].assets) == 1


def test_ambiguous_workbook_sheets_are_not_defaulted_to_first(tmp_path: Path) -> None:
    from openpyxl import Workbook

    source = tmp_path / "ambiguous.xlsx"
    workbook = Workbook()
    workbook.active.title = "Sheet 1"
    workbook.create_sheet("Sheet 2")
    workbook.save(source)
    workbook.close()

    result = AssetScanner().scan((source,))

    assert len(result.assets) == 1
    assert result.assets[0].shape_hint is None
    assert any("未替你任选" in evidence.basis for evidence in result.assets[0].evidence)


def test_bounded_excel_preview_does_not_materialize_unshown_columns(tmp_path: Path) -> None:
    from openpyxl import Workbook

    source = tmp_path / "wide.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["gene_id", *[f"sample_{index}" for index in range(100)]])
    sheet.append(["gene_a", *[index + 1 for index in range(100)]])
    workbook.save(source)
    workbook.close()

    preview = TableReader().read(source, max_rows=40, max_columns=80)
    complete = TableReader().read(source)

    assert preview.column_count == 80
    assert all(len(row) == 80 for row in preview.rows)
    assert preview.is_complete is False
    assert complete.column_count == 101
    assert len(complete.rows[0]) == 101


def test_text_preview_rejects_a_single_pathologically_wide_physical_line(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from rnaseq_helper.readers.table_reader import TableReadLimitError

    source = tmp_path / "wide-row.tsv"
    source.write_text("\t".join("x" for _ in range(101)) + "\n", encoding="utf-8")
    monkeypatch.setattr(TableReader, "MAX_PREVIEW_PHYSICAL_LINE_CHARACTERS", 64)

    with pytest.raises(TableReadLimitError, match="安全预览上限"):
        TableReader().read(source, max_rows=40, max_columns=80)


def test_text_preview_caps_retained_cell_text_and_can_retry_after_simplifying(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from rnaseq_helper.readers.table_reader import TableReadLimitError

    source = tmp_path / "long-annotation.tsv"
    source.write_text("gene_id\tsample_1\tsample_2\n" + "g1\t" + "x" * 32 + "\t2\n", encoding="utf-8")
    monkeypatch.setattr(TableReader, "MAX_PREVIEW_RETAINED_CHARACTERS", 30)

    with pytest.raises(TableReadLimitError, match="预览内容超过安全上限"):
        TableReader().read(source, max_rows=40, max_columns=80)

    monkeypatch.setattr(TableReader, "MAX_PREVIEW_RETAINED_CHARACTERS", 64)
    preview = TableReader().read(source, max_rows=40, max_columns=80)
    assert preview.rows[1][0] == "g1"
    assert preview.is_complete is True


def test_text_preview_translates_oversized_cell_to_recoverable_error(tmp_path: Path) -> None:
    import csv

    from rnaseq_helper.readers.table_reader import TableReadLimitError

    source = tmp_path / "long-cell.tsv"
    source.write_text("gene_id\tsample\ngene_a\t" + "x" * 32 + "\n", encoding="utf-8")
    original_limit = csv.field_size_limit()
    try:
        csv.field_size_limit(16)
        with pytest.raises(TableReadLimitError, match="单个格子的内容超过安全预览上限"):
            TableReader().read(source, max_rows=40, max_columns=80)
    finally:
        csv.field_size_limit(original_limit)
