from __future__ import annotations

import json
import os
from pathlib import Path
from types import SimpleNamespace

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from openpyxl import load_workbook
import pytest
from PySide6.QtWidgets import QApplication

from rnaseq_helper.analysis.service import AnalysisService
from rnaseq_helper.domain.models import DEGConfig
from rnaseq_helper.services.report_service import ReportService
from rnaseq_helper.ui.pages.results_page import ResultsPage


def test_xlsx_export_is_a_minimal_readable_result_table(tmp_path: Path) -> None:
    output = tmp_path / "deg-results.xlsx"
    result = {
        "table": [
            {
                "feature_key": "feature-1",
                "gene_symbol": None,
                "log2FoldChange": 1.25,
                "pvalue": float("nan"),
                "padj": None,
                "tested": False,
                "reason_not_tested": "all_zero",
            },
            {
                "feature_key": "feature-2",
                "gene_symbol": "GENE2",
                "log2FoldChange": -0.5,
                "pvalue": 0.01,
                "padj": 0.02,
                "direction": "down",
            },
        ]
    }

    ReportService._write_xlsx(output, object(), object(), result, "ignored", {})
    AnalysisService._write_result(
        tmp_path, SimpleNamespace(as_dict=lambda: {"plan_id": "test"}), result,
    )
    saved_result = json.loads((tmp_path / "analysis" / "deg-result.json").read_text(encoding="utf-8"))

    workbook = load_workbook(output, data_only=True)
    assert workbook.sheetnames == ["DEG Results"]
    sheet = workbook["DEG Results"]
    assert [cell.value for cell in sheet[1]] == list(ReportService.RESULT_EXPORT_COLUMNS)
    assert sheet.max_column == len(ReportService.RESULT_EXPORT_COLUMNS)
    assert [cell.value for cell in sheet[2]] == ["feature-1", "NA", 1.25, "NA", "NA"]
    assert [cell.value for cell in sheet[3]] == ["feature-2", "GENE2", -0.5, 0.01, 0.02]
    json_rows = saved_result["table"]
    expected = [
        [ReportService._export_result_value(row.get(column)) for column in ReportService.RESULT_EXPORT_COLUMNS]
        for row in json_rows
    ]
    actual = [
        [cell.value for cell in sheet[row_number]]
        for row_number in range(2, sheet.max_row + 1)
    ]
    assert actual == expected
    workbook.close()


def test_html_result_cells_match_xlsx_missing_value_display_and_escape_text() -> None:
    assert ReportService._html_result_value(None) == "NA"
    assert ReportService._html_result_value(float("nan")) == "NA"
    assert ReportService._html_result_value("") == "NA"
    assert ReportService._html_result_value(" <gene&> ") == " &lt;gene&amp;&gt; "


def test_html_report_distinguishes_matrix_columns_from_analyzed_samples(tmp_path: Path) -> None:
    output = tmp_path / "analysis-report.html"
    matrix = SimpleNamespace(
        values=SimpleNamespace(shape=(120, 8)),
        content_hash="matrix-hash",
        source_modality="",
        platform_id="",
        measurement_provenance="",
        validation=SimpleNamespace(issues=()),
    )
    assignments = {
        "C1": "control", "C2": "control", "C3": "control",
        "T1": "test", "T2": "test", "T3": "test",
        "X1": "excluded", "X2": "excluded",
    }
    plan = SimpleNamespace(
        method="DESEQ2",
        method_profile_id="deseq2-count-v1",
        limitations_zh=(),
        warnings=(),
        design={
            "assignments": assignments,
            "replicate_structure": "paired",
            "replicate_structure_confirmed": True,
            "type": "paired",
            "paired_subjects": {
                "C1": "patient-001", "T1": "patient-001",
                "C2": "patient-002", "T2": "patient-002",
                "C3": "patient-003", "T3": "patient-003",
            },
            "batch": {sample: "batch-private-label" for sample in assignments},
        },
        contrast={"control": "control", "test": "test"},
        method_reason_zh="reason",
    )

    ReportService._write_html(output, matrix, plan, {"table": []}, {}, "", {})

    report = output.read_text(encoding="utf-8")
    assert "完整矩阵：120 个特征 × 8 个样本列" in report
    assert "分析纳入样本：6（control 3；test 3）" in report
    assert "对比方向：test − control（正值表示测试组在模型尺度上高于对照组）" in report
    assert "未纳入的样本不参与本次统计" in report
    assert "样本关系：已确认同一对象的配对/重复测量" in report
    assert "C1: control；配对个体 个体1；批次1" in report
    assert "T1: test；配对个体 个体1；批次1" in report
    assert "patient-001" not in report
    assert "batch-private-label" not in report


def test_report_transform_explains_analysis_step_and_unknown_upstream_formula() -> None:
    linear_plan = SimpleNamespace(
        method="LIMMA_ON_LOG2_TPM",
        preprocessing={"transform": "log2(x+1)", "filter_threshold": 1.0},
    )
    assert ReportService._transform_description(linear_plan, {}) == (
        "软件本次对线性输入先按 x≥1.0 过滤，再执行 log2(x+1)；未做额外重新标准化。"
    )

    logged_plan = SimpleNamespace(
        method="LIMMA_ON_LOG2_TPM",
        preprocessing={"transform": "already_log2"},
    )
    assert "未再次取对数" in ReportService._transform_description(logged_plan, {})
    assert "未记录可核实的原文件公式" in ReportService._transform_description(logged_plan, {})
    assert "log2(TPM + 1)" in ReportService._transform_description(
        logged_plan, {"reported_transform": "log2(TPM + 1)"}
    )


def test_results_page_explains_limma_effect_scale_without_changing_xlsx_columns(tmp_path: Path) -> None:
    application = QApplication.instance() or QApplication([])
    matrix_path = tmp_path / "prepared" / "expression_matrix.tsv"
    matrix_path.parent.mkdir(parents=True)
    matrix_path.write_text("feature_key\tS1\n", encoding="utf-8")
    page = ResultsPage()
    page.matrix = SimpleNamespace(values=SimpleNamespace(path=matrix_path))

    page.set_result({"method": "LIMMA_ON_LOG2_TPM", "issues": []})

    assert "当前 log 表达尺度上的模型差异" in page.status_label.text()
    assert list(ReportService.RESULT_EXPORT_COLUMNS) == [
        "feature_key", "gene_symbol", "log2FoldChange", "pvalue", "padj",
    ]
    page.close()
    application.processEvents()


def test_limma_worker_filter_threshold_must_match_confirmed_plan() -> None:
    plan = SimpleNamespace(
        method="LIMMA_ON_LOG2_TPM",
        preprocessing={"transform": "log2(x+1)", "filter_threshold": 2.0},
        config=DEGConfig("control", "test"),
    )
    parameters = {
        "p_adjust_method": "BH",
        "transform": "log2(x+1)",
        "filter_threshold": 1.0,
        "trend": True,
        "robust": True,
        "winsor.tail.p": [0.05, 0.1],
        "proportion": 0.01,
        "stdev.coef.lim": [0.1, 4],
        "legacy_supported": True,
        "legacy": True,
    }

    with pytest.raises(ValueError, match="过滤阈值与分析计划不一致"):
        AnalysisService._validate_parameters(plan, parameters)
