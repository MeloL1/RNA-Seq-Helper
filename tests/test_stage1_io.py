import csv
import hashlib
import json
import tempfile
import unittest
from dataclasses import replace
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook
import xlwt

from rnaseq_helper.preparation.service import PreparationBlocked, PreparationCancelled, PreparationService
from rnaseq_helper.recognition.scanner import AssetScanner
from rnaseq_helper.domain.models import MeasurementType, ObservationUnit, ValueScale
from tests._planning import PlanningService
from rnaseq_helper.readers.geo_reader import GeoSeriesMatrixReader
from rnaseq_helper.readers.table_reader import TableReader
from rnaseq_helper.services.reference_service import ReferenceService
from rnaseq_helper.services.project_store import ProjectStore


FIXTURE = Path(__file__).parent / "fixtures" / "company_counts.tsv"


class Stage1IoTests(unittest.TestCase):
    def test_archive_member_is_bundled_and_project_source_survives_relocation(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "expanded" / "counts.tsv"
            source.parent.mkdir()
            source.write_bytes(FIXTURE.read_bytes())
            asset = replace(
                AssetScanner().scan((source,)).assets[0],
                archive_member_paths=("counts.tsv",),
            )
            project = root / "project"
            service = PreparationService()
            session = service.start(asset, project)
            session = service.apply(session, {
                "measurement_type": MeasurementType.RAW_COUNT,
                "value_scale": ValueScale.LINEAR,
            })
            service.finish(session)

            manifest = json.loads((project / "project.json").read_text(encoding="utf-8"))
            self.assertTrue(manifest["source_path_relative"])
            self.assertEqual(manifest["source_path"], "source-input/counts.tsv")
            self.assertEqual(
                hashlib.sha256((project / manifest["source_path"]).read_bytes()).hexdigest(),
                manifest["source_sha256"],
            )
            initial_warnings = ProjectStore.load(project).warnings
            self.assertFalse(any("项目内保存的来源文件无法访问" in warning for warning in initial_warnings))

            relocated = root / "relocated-project"
            project.rename(relocated)
            snapshot = ProjectStore.load(relocated)
            self.assertFalse(any("来源文件无法访问" in warning for warning in snapshot.warnings))
            self.assertFalse(any("来源文件内容与导入时记录不一致" in warning for warning in snapshot.warnings))

    def test_scan_finds_table_and_hints_shape(self):
        result = AssetScanner().scan((FIXTURE,))
        self.assertEqual(len(result.assets), 1)
        self.assertEqual(result.assets[0].kind, "TABLE")
        self.assertEqual(result.assets[0].shape_hint, (7, 8))

    def test_prepare_requires_unit_and_scale_confirmation(self):
        asset = AssetScanner().scan((FIXTURE,)).assets[0]
        with tempfile.TemporaryDirectory() as temp:
            service = PreparationService()
            session = service.start(asset, Path(temp))
            self.assertEqual(session.recipe.header_row, 2)
            self.assertEqual(session.recipe.sample_columns, (2, 3, 4, 5, 6, 7))
            unknown_summary = service.validate(session)
            self.assertFalse(unknown_summary.passed)
            self.assertTrue({"UNKNOWN_MEASUREMENT", "UNKNOWN_SCALE"}.issubset(
                {issue.code for issue in unknown_summary.issues}
            ))
            with self.assertRaises(PreparationBlocked):
                service.finish(session, summary=unknown_summary)
            self.assertFalse((Path(temp) / "prepared" / "expression_matrix.tsv").exists())
            session = service.apply(session, {"measurement_type": MeasurementType.RAW_COUNT, "value_scale": ValueScale.LINEAR})
            summary = service.validate(session)
            self.assertTrue(summary.passed)
            matrix = service.finish(session)
            self.assertEqual(matrix.values.shape, (4, 6))
            self.assertEqual(matrix.values.path, Path(temp) / "prepared" / "expression_matrix.tsv")
            self.assertTrue(matrix.values.path.is_file())
            self.assertTrue((Path(temp) / "prepared" / "expression_matrix.tsv").exists())
            self.assertTrue((Path(temp) / "preparation-recipe.json").exists())
            self.assertEqual(list(Path(temp).glob(".rnaseq-helper-preparing-*")), [])
            project_manifest = json.loads((Path(temp) / "project.json").read_text(encoding="utf-8"))
            self.assertEqual(project_manifest["source_sha256"], hashlib.sha256(FIXTURE.read_bytes()).hexdigest())
            feature_metadata = (Path(temp) / "prepared" / "feature_metadata.tsv").read_text(encoding="utf-8")
            self.assertIn("gene_symbol", feature_metadata.splitlines()[0])

    def test_count_above_r_integer_limit_is_blocked_during_preparation(self):
        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / "large-counts.tsv"
            source.write_text(
                "gene_id\tC1\tC2\tT1\tT2\n"
                "g1\t2147483648\t10\t12\t14\n",
                encoding="utf-8",
            )
            asset = AssetScanner().scan((source,)).assets[0]
            service = PreparationService()
            session = service.start(asset, Path(temp) / "project")
            session = service.apply(session, {
                "measurement_type": MeasurementType.RAW_COUNT,
                "value_scale": ValueScale.LINEAR,
            })

            summary = service.validate(session)

            self.assertFalse(summary.passed)
            self.assertTrue(any(
                issue.code == "RAW_COUNT_EXCEEDS_INTEGER_RANGE" for issue in summary.issues
            ))

    def test_negative_and_fractional_counts_are_blocked_during_preparation(self):
        for value in ("-1", "1.5"):
            with self.subTest(value=value), tempfile.TemporaryDirectory() as temp:
                source = Path(temp) / "invalid-counts.tsv"
                source.write_text(
                    f"gene_id\tC1\tC2\tT1\tT2\ng1\t10\t11\t{value}\t14\n",
                    encoding="utf-8",
                )
                asset = AssetScanner().scan((source,)).assets[0]
                service = PreparationService()
                session = service.start(asset, Path(temp) / "project")
                session = service.apply(session, {
                    "measurement_type": MeasurementType.RAW_COUNT,
                    "value_scale": ValueScale.LINEAR,
                })

                summary = service.validate(session)

                self.assertFalse(summary.passed)
                self.assertTrue(any(issue.code == "INVALID_RAW_COUNT" for issue in summary.issues))

    def test_negative_linear_tpm_is_blocked_without_clipping(self):
        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / "negative-tpm.tsv"
            source.write_text(
                "gene_id\tC1\tC2\tT1\tT2\n"
                "g1\t10\t11\t-0.5\t14\n",
                encoding="utf-8",
            )
            asset = AssetScanner().scan((source,)).assets[0]
            service = PreparationService()
            session = service.start(asset, Path(temp) / "project")
            session = service.apply(session, {
                "measurement_type": MeasurementType.TPM,
                "value_scale": ValueScale.LINEAR,
            })

            summary = service.validate(session)

            self.assertFalse(summary.passed)
            self.assertTrue(any(issue.code == "NEGATIVE_LINEAR_VALUE" for issue in summary.issues))
            with self.assertRaises(PreparationBlocked):
                service.finish(session, summary=summary)
            self.assertFalse((Path(temp) / "project" / "prepared" / "expression_matrix.tsv").exists())

    def test_nonfinite_expression_values_are_blocked_during_preparation(self):
        for value in ("NaN", "Inf", "-Inf"):
            with self.subTest(value=value), tempfile.TemporaryDirectory() as temp:
                source = Path(temp) / "nonfinite-expression.tsv"
                source.write_text(
                    f"gene_id\tC1\tC2\tT1\tT2\ng1\t10\t11\t{value}\t14\n",
                    encoding="utf-8",
                )
                asset = AssetScanner().scan((source,)).assets[0]
                service = PreparationService()
                session = service.start(asset, Path(temp) / "project")
                session = service.apply(session, {
                    "measurement_type": MeasurementType.TPM,
                    "value_scale": ValueScale.LINEAR,
                })

                summary = service.validate(session)

                self.assertFalse(summary.passed)
                self.assertTrue(any(issue.code == "NONFINITE_VALUE" for issue in summary.issues))
                with self.assertRaises(PreparationBlocked):
                    service.finish(session, summary=summary)
                self.assertFalse((Path(temp) / "project" / "prepared" / "expression_matrix.tsv").exists())

    def test_cell_observation_hint_survives_preparation_and_blocks_bulk_analysis(self):
        asset = AssetScanner().scan((FIXTURE,)).assets[0]
        asset = replace(asset, observation_unit_hint="CELL")
        with tempfile.TemporaryDirectory() as temp:
            project = Path(temp) / "project"
            service = PreparationService()
            session = service.start(asset, project)
            session = service.apply(session, {
                "measurement_type": MeasurementType.RAW_COUNT,
                "value_scale": ValueScale.LINEAR,
            })

            matrix = service.finish(session)
            manifest = json.loads((project / "project.json").read_text(encoding="utf-8"))
            with matrix.values.path.open(encoding="utf-8", newline="") as handle:
                sample_ids = next(csv.reader(handle, delimiter="\t"))[1:]
            assignments = {
                sample_id: "control" if index < len(sample_ids) // 2 else "test"
                for index, sample_id in enumerate(sample_ids)
            }
            plan = PlanningService().build(matrix, assignments)

            self.assertEqual(matrix.observation_unit, ObservationUnit.CELL)
            self.assertEqual(manifest["observation_unit"], "CELL")
            self.assertTrue(any(issue.code == "CELL_NOT_BULK" for issue in plan.blockers))

    def test_unlabeled_expression_table_does_not_assume_bulk_observation_unit(self):
        asset = AssetScanner().scan((FIXTURE,)).assets[0]
        with tempfile.TemporaryDirectory() as temp:
            project = Path(temp) / "project"
            service = PreparationService()
            session = service.start(asset, project)
            session = service.apply(session, {
                "measurement_type": MeasurementType.RAW_COUNT,
                "value_scale": ValueScale.LINEAR,
            })

            matrix = service.finish(session)
            manifest = json.loads((project / "project.json").read_text(encoding="utf-8"))

            self.assertEqual(matrix.observation_unit, ObservationUnit.UNKNOWN)
            self.assertEqual(manifest["observation_unit"], "UNKNOWN")

    def test_original_file_is_unchanged_and_duplicate_ids_block(self):
        before = hashlib.sha256(FIXTURE.read_bytes()).hexdigest()
        asset = AssetScanner().scan((FIXTURE,)).assets[0]
        with tempfile.TemporaryDirectory() as temp:
            service = PreparationService()
            session = service.start(asset, Path(temp))
            rows = list(session.raw_table.rows)
            rows[4] = (rows[3][0], *rows[4][1:])
            session = replace(session, raw_table=replace(session.raw_table, rows=tuple(rows)))
            session = service.apply(session, {"measurement_type": "RAW_COUNT", "value_scale": "LINEAR"})
            with self.assertRaises(PreparationBlocked):
                service.finish(session)
        self.assertEqual(hashlib.sha256(FIXTURE.read_bytes()).hexdigest(), before)

    def test_missing_source_during_finish_does_not_leave_partial_project(self):
        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / "counts.tsv"
            source.write_bytes(FIXTURE.read_bytes())
            project = Path(temp) / "project"
            service = PreparationService()
            asset = AssetScanner().scan((source,)).assets[0]
            session = service.start(asset, project)
            session = service.apply(session, {
                "measurement_type": MeasurementType.RAW_COUNT,
                "value_scale": ValueScale.LINEAR,
            })
            summary = service.validate(session)
            self.assertTrue(summary.passed)
            source.unlink()

            with self.assertRaises(FileNotFoundError):
                service.finish(session, summary=summary)

            self.assertFalse((project / "project.json").exists())
            self.assertFalse((project / "prepared" / "expression_matrix.tsv").exists())

    def test_source_fingerprint_reports_progress_and_honors_cancellation(self):
        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / "large-source.tsv"
            source.write_bytes(b"x" * (9 * 1024 * 1024))
            events = []
            cancel = False

            def report_progress(completed, total):
                nonlocal cancel
                events.append((completed, total))
                if completed >= 8 * 1024 * 1024:
                    cancel = True

            with self.assertRaises(PreparationCancelled):
                PreparationService.fingerprint_source(
                    source,
                    progress=report_progress,
                    cancelled=lambda: cancel,
                )

            self.assertIn((0, source.stat().st_size), events)
            self.assertIn((8 * 1024 * 1024, source.stat().st_size), events)

    def test_source_fingerprint_rejects_file_changed_during_hash(self):
        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / "changing-source.tsv"
            source.write_bytes(b"initial")

            def mutate_after_stat(completed, total):
                if completed == 0:
                    source.write_bytes(b"changed while hashing")

            with self.assertRaisesRegex(ValueError, "来源文件在校验期间发生变化"):
                PreparationService.fingerprint_source(source, progress=mutate_after_stat)

    def test_source_change_after_complete_preview_requires_reimport(self):
        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / "counts.tsv"
            source.write_bytes(FIXTURE.read_bytes())
            asset = AssetScanner().scan((source,)).assets[0]
            service = PreparationService()
            session = service.start(asset, Path(temp) / "project", preview_only=True)
            self.assertTrue(session.raw_table.is_complete)
            source.write_bytes(source.read_bytes() + b"# changed after preview\n")

            with self.assertRaisesRegex(ValueError, "来源文件在预览后发生变化"):
                service.complete_read(session)

    def test_source_change_after_fingerprint_does_not_write_project(self):
        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / "counts.tsv"
            source.write_bytes(FIXTURE.read_bytes())
            project = Path(temp) / "project"
            service = PreparationService()
            asset = AssetScanner().scan((source,)).assets[0]
            session = service.start(asset, project)
            session = service.apply(session, {
                "measurement_type": MeasurementType.RAW_COUNT,
                "value_scale": ValueScale.LINEAR,
            })
            summary = service.validate(session)
            source_hash = service.fingerprint_source(source, expected_signature=session.source_signature)
            source.write_bytes(source.read_bytes() + b"# changed before save\n")

            with self.assertRaisesRegex(ValueError, "来源文件在导入后发生变化"):
                service.finish(session, summary=summary, source_sha256=source_hash)

            self.assertFalse((project / "prepared" / "expression_matrix.tsv").exists())

    def test_local_reference_is_recorded_without_changing_matrix_rows(self):
        with tempfile.TemporaryDirectory() as temp, patch.object(
            ReferenceService, "CACHE_DIR", Path(temp) / "reference-cache",
        ):
            source = Path(temp) / "counts.tsv"
            source.write_text(
                "gene_id\tCtrl_1\tCtrl_2\tTreat_1\tTreat_2\n"
                "g1\t100\t110\t220\t230\n"
                "g2\t50\t52\t55\t53\n"
                "g3\t80\t82\t160\t165\n"
                "g4\t10\t12\t20\t22\n",
                encoding="utf-8",
            )
            asset = AssetScanner().scan((source,)).assets[0]
            reference = Path(temp) / "reference.tsv"
            reference.write_text(
                "feature_key\tgene_symbol\n"
                "g1\tGENE1\n"
                "g2\tGENE2\n"
                "g2\tGENE2_ALT\n",
                encoding="utf-8",
            )
            service = PreparationService()
            session = service.start(asset, Path(temp) / "project")
            session = replace(session, reference_path=reference)
            session = service.apply(session, {"measurement_type": MeasurementType.RAW_COUNT, "value_scale": ValueScale.LINEAR})
            matrix = service.finish(session)
            feature_rows = (Path(temp) / "project" / "prepared" / "feature_metadata.tsv").read_text(encoding="utf-8").splitlines()
            assert feature_rows[0].endswith("\tgene_symbol")
            assert feature_rows[1].endswith("\tGENE1")
            assert feature_rows[2].endswith("\tGENE2;GENE2_ALT")
            assert len((Path(temp) / "project" / "prepared" / "expression_matrix.tsv").read_text(encoding="utf-8").splitlines()) == 5
            assert (Path(temp) / "project" / "references" / "reference-result.json").exists()
            assert matrix.feature_metadata.shape == (4, 3)
            self.assertEqual(len(list((Path(temp) / "reference-cache").glob("*.json"))), 1)

    def test_reference_resolution_failure_does_not_leave_partial_project(self):
        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / "counts.tsv"
            source.write_bytes(FIXTURE.read_bytes())
            reference = Path(temp) / "reference.tsv"
            reference.write_text("feature_key\tgene_symbol\ngene-a\tA\n", encoding="utf-8")
            project = Path(temp) / "project"
            service = PreparationService()
            asset = AssetScanner().scan((source,)).assets[0]
            session = service.start(asset, project)
            session = replace(session, reference_path=reference)
            session = service.apply(session, {
                "measurement_type": MeasurementType.RAW_COUNT,
                "value_scale": ValueScale.LINEAR,
            })

            with patch.object(ReferenceService, "resolve", side_effect=OSError("test reference read failure")):
                with self.assertRaisesRegex(OSError, "test reference read failure"):
                    service.finish(session)

            self.assertFalse(project.exists())
            self.assertEqual(list(Path(temp).glob(".rnaseq-helper-preparing-*")), [])

    def test_preparation_refuses_to_replace_existing_project_files(self):
        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / "counts.tsv"
            source.write_bytes(FIXTURE.read_bytes())
            project = Path(temp) / "project"
            project.mkdir()
            existing = project / "keep.txt"
            existing.write_text("existing project data", encoding="utf-8")
            service = PreparationService()
            asset = AssetScanner().scan((source,)).assets[0]
            session = service.start(asset, project)
            session = service.apply(session, {
                "measurement_type": MeasurementType.RAW_COUNT,
                "value_scale": ValueScale.LINEAR,
            })

            with self.assertRaisesRegex(FileExistsError, "避免覆盖旧项目"):
                service.finish(session)

            self.assertEqual(existing.read_text(encoding="utf-8"), "existing project data")
            self.assertFalse((project / "prepared").exists())

    def test_cancelled_preparation_validation_cleans_staging_directory(self):
        with tempfile.TemporaryDirectory() as temp:
            project = Path(temp) / "project"
            asset = AssetScanner().scan((FIXTURE,)).assets[0]
            service = PreparationService()
            session = service.start(asset, project)
            session = service.apply(session, {
                "measurement_type": MeasurementType.RAW_COUNT,
                "value_scale": ValueScale.LINEAR,
            })

            with self.assertRaises(PreparationCancelled):
                service.finish(session, cancelled=lambda: True)

            self.assertFalse(project.exists())
            self.assertEqual(list(Path(temp).glob(".rnaseq-helper-preparing-*")), [])

    def test_cancel_after_materialization_cleans_staging_before_matrix_write(self):
        with tempfile.TemporaryDirectory() as temp:
            project = Path(temp) / "project"
            asset = AssetScanner().scan((FIXTURE,)).assets[0]
            service = PreparationService()
            session = service.start(asset, project)
            session = service.apply(session, {
                "measurement_type": MeasurementType.RAW_COUNT,
                "value_scale": ValueScale.LINEAR,
            })
            summary = service.validate(session)
            source_hash = service.fingerprint_source(FIXTURE)
            original_materialize = service._materialize
            state = {"cancelled": False}

            def materialize_then_cancel(*args, **kwargs):
                result = original_materialize(*args, **kwargs)
                state["cancelled"] = True
                return result

            with patch.object(service, "_materialize", side_effect=materialize_then_cancel):
                with self.assertRaises(PreparationCancelled):
                    service.finish(
                        session,
                        summary=summary,
                        source_sha256=source_hash,
                        cancelled=lambda: state["cancelled"],
                    )

            self.assertFalse(project.exists())
            self.assertEqual(list(Path(temp).glob(".rnaseq-helper-preparing-*")), [])

    def test_excel_sheet_and_formula_cache_are_read_without_execution(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "input.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "说明"
            sheet.append(["note", "not expression"])
            data = workbook.create_sheet("表达")
            data.append(["gene_id", "C1", "C2"])
            data.append(["g1", 1, 2])
            data["C3"] = "=1+1"
            workbook.save(path)
            reader = TableReader()
            self.assertEqual(reader.list_sheets(path), ("说明", "表达"))
            table = reader.read(path, "表达")
            self.assertEqual(table.rows[1], ("g1", "1", "2"))
            self.assertIn("C3", table.formula_cells)
            asset = AssetScanner().scan((path,)).assets[0]
            service = PreparationService()
            session = service.start(asset, Path(temp), "表达")
            session = service.apply(session, {"measurement_type": "RAW_COUNT", "value_scale": "LINEAR"})
            summary = service.validate(session)
            self.assertTrue(any(issue.code == "FORMULA_WITHOUT_CACHE" for issue in summary.issues))
            self.assertFalse(summary.passed)

    def test_legacy_xls_is_read_and_scanned(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "legacy.xls"
            workbook = xlwt.Workbook()
            notes = workbook.add_sheet("说明")
            notes.write(0, 0, "这是说明表")
            data = workbook.add_sheet("表达")
            for column, value in enumerate(("gene_id", "C1", "C2")):
                data.write(0, column, value)
            data.write(1, 0, "g1")
            data.write(1, 1, 1)
            data.write(1, 2, 2)
            data.write(2, 0, "g2")
            data.write(2, 1, 3)
            data.write(2, 2, xlwt.Formula("1+1"))
            workbook.save(str(path))

            reader = TableReader()
            self.assertEqual(reader.list_sheets(path), ("说明", "表达"))
            table = reader.read(path, "表达")
            self.assertEqual(table.rows[1], ("g1", "1.0", "2.0"))
            asset = AssetScanner().scan((path,)).assets[0]
            self.assertEqual(asset.kind, "TABLE")
            self.assertFalse(GeoSeriesMatrixReader.looks_like(path))

    def test_vendor_text_matrix_named_xls_ignores_unit_in_filename(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "gene.TPM.matrix.annotation.xls"
            path.write_text(
                "Gene id\tGene name\tBio_type\tDZ_1\tCA_1\tYDW_1\n"
                "ENSMUSG0001\tGeneA\tprotein_coding\t1.0\t2.0\t3.0\n"
                "ENSMUSG0002\tGeneB\tprotein_coding\t4.0\t5.0\t6.0\n",
                encoding="utf-8",
            )
            reader = TableReader()
            self.assertEqual(reader.list_sheets(path), (None,))
            table = reader.read(path)
            self.assertEqual(table.delimiter, "\t")
            self.assertEqual(table.rows[0][:4], ("Gene id", "Gene name", "Bio_type", "DZ_1"))
            self.assertEqual(table.rows[1][4], "2.0")
            asset = AssetScanner().scan((path,)).assets[0]
            self.assertEqual(asset.kind, "TABLE")
            self.assertEqual(asset.measurement_hint, MeasurementType.UNKNOWN)
            self.assertEqual(asset.scale_hint, ValueScale.UNKNOWN)


if __name__ == "__main__":
    unittest.main()
