from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any

from PySide6.QtCore import QAbstractTableModel, QModelIndex, QSortFilterProxyModel, QThread, QUrl, Qt, Signal, Slot
from PySide6.QtGui import QDesktopServices, QPixmap, QGuiApplication
from PySide6.QtWidgets import (
    QAbstractItemView, QCheckBox, QDialog, QDialogButtonBox, QDoubleSpinBox,
    QFileDialog, QFormLayout, QFrame, QGroupBox, QHBoxLayout, QLabel,
    QLineEdit, QListWidget, QMessageBox, QPushButton, QRadioButton,
    QComboBox, QSplitter, QTabWidget, QTableView, QTableWidget, QTableWidgetItem,
    QVBoxLayout, QWidget,
)

from rnaseq_helper.analysis.result_filter import passes_result_filter, summarize_result
from rnaseq_helper.domain.models import ExpressionMatrix, JobContext, ResultFilterConfig
from rnaseq_helper.ui.workers.result_filter_worker import ResultFilterWorker
from rnaseq_helper.ui.workers.annotation_worker import AnnotationWorker


class _GeneTableModel(QAbstractTableModel):
    HEADERS = ("Gene ID", "Gene Symbol", "log2FoldChange", "原始 P 值", "FDR", "当前筛选", "分析状态/未检测原因")
    REASONS = {
        "below_expression_filter": "未检测：低表达过滤",
        "all_zero": "未检测：全零表达",
        "independent_filtering": "已检测；独立过滤后未提供 FDR",
        "cooks_outlier": "未检测：Cook 异常值",
        "not_estimable": "未检测：模型不可估计",
        "degenerate_residual_variance": "未检测：残差方差退化",
        "invalid_or_degenerate_variance": "未检测：方差无效或退化",
    }
    SORT_ROLE = Qt.ItemDataRole.UserRole + 1

    def __init__(self, parent=None):
        super().__init__(parent)
        self.rows: list[dict[str, Any]] = []
        self.symbols: dict[str, str] = {}
        self.filter_config = ResultFilterConfig()

    def set_data(self, rows, symbols, config):
        self.beginResetModel()
        self.rows = [row for row in rows if isinstance(row, dict)]
        self.symbols = dict(symbols)
        self.filter_config = config
        self.endResetModel()

    def rowCount(self, parent=QModelIndex()):
        return 0 if parent.isValid() else len(self.rows)

    def columnCount(self, parent=QModelIndex()):
        return 0 if parent.isValid() else len(self.HEADERS)

    def headerData(self, section, orientation, role=Qt.ItemDataRole.DisplayRole):
        if role == Qt.ItemDataRole.DisplayRole and orientation == Qt.Orientation.Horizontal and 0 <= section < len(self.HEADERS):
            return self.HEADERS[section]
        return None

    def data(self, index, role=Qt.ItemDataRole.DisplayRole):
        if not index.isValid() or not 0 <= index.row() < len(self.rows):
            return None
        row = self.rows[index.row()]
        key = str(row.get("feature_key") or row.get("original_id") or "")
        column = index.column()
        tested = row.get("tested") is True
        reason = str(row.get("reason_not_tested") or "")
        values = (
            key,
            self.symbols.get(key, str(row.get("gene_symbol") or "")),
            row.get("log2FoldChange"), row.get("pvalue"), row.get("padj"),
            ("显著" if passes_result_filter(row, self.filter_config) else "未达到当前条件") if tested else "未检测",
            (self.REASONS.get(reason, "已检测") if tested else self.REASONS.get(reason, f"未检测：{reason or '原因未记录'}")),
        )
        value = values[column]
        if role == self.SORT_ROLE:
            if column in (2, 3, 4):
                return float(value) if isinstance(value, (int, float)) else float("inf")
            return str(value).casefold()
        if role == Qt.ItemDataRole.DisplayRole:
            return f"{value:.5g}" if isinstance(value, (int, float)) else value
        if role == Qt.ItemDataRole.UserRole:
            return row
        return None


class _GeneFilterProxy(QSortFilterProxyModel):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.query = ""
        self.significant_only = True
        self.hide_untested = True
        self.config = ResultFilterConfig()
        self.setFilterCaseSensitivity(Qt.CaseSensitivity.CaseInsensitive)
        self.setSortRole(_GeneTableModel.SORT_ROLE)
        self.setDynamicSortFilter(True)

    def filterAcceptsRow(self, source_row, source_parent):
        model = self.sourceModel()
        if model is None:
            return False
        row = model.rows[source_row]
        if self.hide_untested and row.get("tested") is not True:
            return False
        key = str(row.get("feature_key") or row.get("original_id") or "")
        symbol = model.symbols.get(key, str(row.get("gene_symbol") or ""))
        if self.query and self.query.casefold() not in f"{key} {symbol}".casefold():
            return False
        return not self.significant_only or passes_result_filter(row, self.config)


class _AnnotationPreviewDialog(QDialog):
    """Explicitly select annotation columns and preview mapping coverage before applying."""

    def __init__(self, headers: list[str], rows: list[list[str]], feature_ids: set[str], parent=None):
        super().__init__(parent)
        self.setWindowTitle("核对 Gene ID 与 Gene Symbol 映射")
        self.resize(760, 560)
        self.rows = rows
        self.feature_ids = feature_ids
        self.mapping: dict[str, str] = {}
        layout = QVBoxLayout(self)
        form = QFormLayout()
        self.id_column = QComboBox()
        self.symbol_column = QComboBox()
        self.id_column.addItems(headers)
        self.symbol_column.addItems(headers)
        id_guess = next((i for i, item in enumerate(headers) if item.casefold() in {"feature_key", "gene_id", "feature_id", "seq_id", "id"}), 0)
        symbol_guess = next((i for i, item in enumerate(headers) if item.casefold() in {"gene_symbol", "symbol", "gene name", "gene_name"}), min(1, len(headers) - 1))
        self.id_column.setCurrentIndex(id_guess)
        self.symbol_column.setCurrentIndex(symbol_guess)
        self.id_column.currentIndexChanged.connect(self._refresh)
        self.symbol_column.currentIndexChanged.connect(self._refresh)
        form.addRow("Gene ID 列", self.id_column)
        form.addRow("Gene Symbol 列", self.symbol_column)
        layout.addLayout(form)
        self.summary = QLabel("")
        self.summary.setWordWrap(True)
        layout.addWidget(self.summary)
        self.preview = QTableWidget(0, 2)
        self.preview.setHorizontalHeaderLabels(("Gene ID", "Gene Symbol"))
        self.preview.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        layout.addWidget(self.preview, 1)
        buttons = QDialogButtonBox(QDialogButtonBox.StandardButton.Cancel | QDialogButtonBox.StandardButton.Ok)
        buttons.button(QDialogButtonBox.StandardButton.Ok).setText("确认应用注释")
        buttons.button(QDialogButtonBox.StandardButton.Cancel).setText("取消")
        buttons.accepted.connect(self.accept)
        buttons.rejected.connect(self.reject)
        layout.addWidget(buttons)
        self._refresh()

    def _refresh(self, *_args):
        id_index = self.id_column.currentIndex()
        symbol_index = self.symbol_column.currentIndex()
        if id_index == symbol_index:
            self.mapping = {}
            self.summary.setText("Gene ID 与 Gene Symbol 必须选择不同的列。")
            return
        options: dict[str, set[str]] = {}
        for row in self.rows:
            if max(id_index, symbol_index) >= len(row):
                continue
            key, symbol = row[id_index].strip(), row[symbol_index].strip()
            if key and symbol:
                options.setdefault(key, set()).add(symbol)
        self.mapping = {key: next(iter(symbols)) for key, symbols in options.items() if len(symbols) == 1}
        ambiguous = sum(len(symbols) > 1 for symbols in options.values())
        matched = len(self.feature_ids.intersection(self.mapping))
        unmatched = len(self.feature_ids - self.mapping.keys())
        self.summary.setText(
            f"匹配 {matched:,} 个当前特征；未匹配 {unmatched:,} 个；"
            f"注释文件中有歧义的 ID {ambiguous:,} 个（歧义项不会自动采用）。"
        )
        self.preview.setRowCount(min(12, len(self.rows)))
        for index, row in enumerate(self.rows[:12]):
            for col, source_col in enumerate((id_index, symbol_index)):
                text = row[source_col] if source_col < len(row) else ""
                self.preview.setItem(index, col, QTableWidgetItem(text))


def _read_annotation_table(path: Path) -> tuple[list[str], list[list[str]]]:
    suffix = path.suffix.casefold()
    if suffix in {".csv", ".tsv", ".txt"}:
        delimiter = "," if suffix == ".csv" else "\t"
        with path.open(encoding="utf-8-sig", newline="") as handle:
            reader = csv.reader(handle, delimiter=delimiter)
            rows = list(reader)
        if len(rows) < 2 or not rows[0]:
            raise ValueError("注释表需要包含表头和至少一行数据。")
        return [str(value).strip() for value in rows[0]], [[str(value).strip() for value in row] for row in rows[1:]]
    if suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            worksheet = workbook.active
            values = worksheet.iter_rows(values_only=True)
            header = next(values, None)
            if not header:
                raise ValueError("注释表没有表头。")
            headers = [str(value or "").strip() for value in header]
            rows = [[str(value or "").strip() for value in row] for row in values]
            return headers, rows
        finally:
            workbook.close()
    raise ValueError("支持的注释文件格式：CSV、TSV、TXT、XLSX、XLSM。")


class ResultsPage(QWidget):
    """分析结果工作台：筛选、图表、差异基因、表达分布与历史记录。"""

    _XLSX_HINT = "统计结果表保留 feature_key、gene_symbol、log2FoldChange、pvalue、padj 五列。"
    regroupRequested = Signal()
    _CHARTS = (
        ("volcano", "火山图", "横轴为 log2FoldChange，纵轴为 −log10(P)；点的颜色按当前结果筛选阈值标示。"),
        ("top_heatmap", "差异基因热图", "显示当前筛选下优先展示的差异特征；颜色为分析使用的变换后表达值。"),
        ("pca", "样本 PCA", "展示样本整体表达变异的主成分投影；组别来自已确认的分析设计。"),
        ("sample_correlation", "样本相关性", "展示样本间表达谱相关性，用于描述性检查，不替代实验设计核验。"),
        ("pvalue_histogram", "P 值分布", "展示被检验特征的原始 P 值分布；不用于单独判定差异。"),
    )

    def __init__(self, parent: QWidget | None = None) -> None:
        super().__init__(parent)
        self.matrix: ExpressionMatrix | None = None
        self._report_path: Path | None = None
        self._result: dict[str, Any] | None = None
        self._plan_data: dict[str, Any] | None = None
        self._filter_threads: list[tuple[QThread, ResultFilterWorker, JobContext]] = []
        self._annotation_threads: list[tuple[QThread, AnnotationWorker, JobContext]] = []
        self._filter_is_applicable = False
        self._gene_index: dict[str, int] = {}
        self._annotations: dict[str, str] = {}
        self._gene_model = _GeneTableModel(self)
        self._gene_proxy = _GeneFilterProxy(self)
        self._gene_proxy.setSourceModel(self._gene_model)
        self._sample_detail_mode = False

        layout = QVBoxLayout(self)
        layout.setContentsMargins(18, 12, 18, 12)
        layout.setSpacing(10)
        title = QLabel("结果工作台")
        title.setObjectName("pageTitle")
        layout.addWidget(title)
        self.status_label = QLabel("完成分析后可查看图表、差异基因和项目记录。")
        self.status_label.setWordWrap(True)
        layout.addWidget(self.status_label)

        filters = QGroupBox("结果筛选")
        filter_layout = QVBoxLayout(filters)
        row = QHBoxLayout()
        self.fdr_radio = QRadioButton("FDR（校正后 P 值）")
        self.raw_p_radio = QRadioButton("原始 P 值")
        self.fdr_radio.setChecked(True)
        row.addWidget(self.fdr_radio)
        row.addWidget(self.raw_p_radio)
        row.addWidget(QLabel("P 阈值"))
        self.p_threshold = QDoubleSpinBox()
        self.p_threshold.setRange(0, 1)
        self.p_threshold.setDecimals(4)
        self.p_threshold.setSingleStep(.01)
        self.p_threshold.setValue(.05)
        row.addWidget(self.p_threshold)
        row.addWidget(QLabel("|log2FC| 阈值"))
        self.lfc_threshold = QDoubleSpinBox()
        self.lfc_threshold.setRange(0, 1_000_000)
        self.lfc_threshold.setDecimals(4)
        self.lfc_threshold.setSingleStep(.25)
        self.lfc_threshold.setValue(1)
        row.addWidget(self.lfc_threshold)
        row.addStretch(1)
        filter_layout.addLayout(row)
        self.filter_summary = QLabel("等待分析结果")
        self.filter_summary.setObjectName("resultSummary")
        self.filter_summary.setWordWrap(True)
        filter_layout.addWidget(self.filter_summary)
        self.filter_caveat = QLabel("筛选只改变展示，不改变模型或原始统计列。")
        self.filter_caveat.setWordWrap(True)
        filter_layout.addWidget(self.filter_caveat)
        actions = QHBoxLayout()
        self.apply_filter_button = QPushButton("应用筛选并刷新图表/报告")
        self.apply_filter_button.setEnabled(False)
        self.apply_filter_button.clicked.connect(self.apply_filter)
        actions.addWidget(self.apply_filter_button)
        self.filter_progress = QLabel("")
        actions.addWidget(self.filter_progress, 1)
        filter_layout.addLayout(actions)
        layout.addWidget(filters)
        for widget in (self.fdr_radio, self.raw_p_radio):
            widget.toggled.connect(self._preview_filter)
        self.p_threshold.valueChanged.connect(self._preview_filter)
        self.lfc_threshold.valueChanged.connect(self._preview_filter)

        self.tabs = QTabWidget()
        layout.addWidget(self.tabs, 1)
        self._build_charts_tab()
        self._build_genes_tab()
        self._build_information_tab()

        delivery = QHBoxLayout()
        self.open_report_button = QPushButton("打开离线 HTML 报告")
        self.open_report_button.setEnabled(False)
        self.open_report_button.clicked.connect(self.open_report)
        delivery.addWidget(self.open_report_button)
        self.open_project_button = QPushButton("打开项目文件夹")
        self.open_project_button.setEnabled(False)
        self.open_project_button.clicked.connect(self.open_project)
        delivery.addWidget(self.open_project_button)
        self.regroup_button = QPushButton("修改分组并重新分析")
        self.regroup_button.clicked.connect(self.regroupRequested.emit)
        delivery.addWidget(self.regroup_button)
        delivery.addStretch(1)
        layout.addLayout(delivery)

    def _build_charts_tab(self) -> None:
        page = QWidget()
        outer = QVBoxLayout(page)
        self.chart_tabs = QTabWidget()
        outer.addWidget(self.chart_tabs, 1)
        self._chart_labels: dict[str, QLabel] = {}
        self._chart_paths: dict[str, Path] = {}
        for key, title, help_text in self._CHARTS:
            pane = QWidget()
            pane_layout = QVBoxLayout(pane)
            help_label = QLabel(help_text)
            help_label.setWordWrap(True)
            pane_layout.addWidget(help_label)
            image = QLabel("完成分析后显示图表")
            image.setAlignment(Qt.AlignmentFlag.AlignCenter)
            image.setMinimumHeight(320)
            image.setScaledContents(False)
            pane_layout.addWidget(image, 1)
            chart_actions = QHBoxLayout()
            help_button = QPushButton("图表说明")
            help_button.clicked.connect(lambda _checked=False, caption=title, explanation=help_text: QMessageBox.information(self, caption, explanation))
            chart_actions.addWidget(help_button)
            zoom_button = QPushButton("放大查看")
            zoom_button.clicked.connect(lambda _checked=False, chart=key: self._zoom_chart(chart))
            chart_actions.addWidget(zoom_button)
            open_button = QPushButton("打开原图")
            open_button.clicked.connect(lambda _checked=False, chart=key: self._open_chart(chart))
            chart_actions.addWidget(open_button)
            chart_actions.addStretch(1)
            pane_layout.addLayout(chart_actions)
            self._chart_labels[key] = image
            self.chart_tabs.addTab(pane, title)
        self.tabs.addTab(page, "图表分析")

    def _build_genes_tab(self) -> None:
        page = QWidget()
        layout = QVBoxLayout(page)
        controls = QHBoxLayout()
        self.gene_search = QLineEdit()
        self.gene_search.setPlaceholderText("搜索基因编号或 gene symbol")
        self.gene_search.setClearButtonEnabled(True)
        self.gene_search.textChanged.connect(self._set_gene_query)
        controls.addWidget(self.gene_search, 1)
        self.significant_only = QCheckBox("只显示符合当前筛选条件的基因")
        self.significant_only.setChecked(True)
        self.significant_only.toggled.connect(self._set_significant_only)
        controls.addWidget(self.significant_only)
        self.hide_untested = QCheckBox("隐藏未检测特征")
        self.hide_untested.setChecked(True)
        self.hide_untested.setToolTip("暂时隐藏低表达、全零或模型无法估计、没有可用检验结果的特征；不会删除或修改分析结果。取消勾选可查看这些行及原因。")
        self.hide_untested.toggled.connect(self._set_hide_untested)
        controls.addWidget(self.hide_untested)
        self.copy_ids_button = QPushButton("复制 Gene ID")
        self.copy_ids_button.clicked.connect(lambda: self._copy_gene_column(0))
        controls.addWidget(self.copy_ids_button)
        self.copy_symbols_button = QPushButton("复制 Gene Symbol")
        self.copy_symbols_button.clicked.connect(lambda: self._copy_gene_column(1))
        controls.addWidget(self.copy_symbols_button)
        self.annotation_button = QPushButton("导入 gene symbol 映射…")
        self.annotation_button.clicked.connect(self.import_annotations)
        controls.addWidget(self.annotation_button)
        self.export_button = QPushButton("导出当前筛选结果…")
        self.export_button.clicked.connect(self.export_filtered)
        controls.addWidget(self.export_button)
        layout.addLayout(controls)
        splitter = QSplitter(Qt.Orientation.Horizontal)
        self.gene_table = QTableView()
        self.gene_table.setModel(self._gene_proxy)
        self.gene_table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.gene_table.setSelectionMode(QAbstractItemView.SelectionMode.ExtendedSelection)
        self.gene_table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.gene_table.setSortingEnabled(True)
        self.gene_table.horizontalHeader().setStretchLastSection(True)
        self.gene_table.setColumnWidth(0, 175)
        self.gene_table.setColumnWidth(1, 150)
        self.gene_table.setColumnWidth(2, 145)
        self.gene_table.setColumnWidth(3, 115)
        self.gene_table.setColumnWidth(4, 115)
        self.gene_table.setColumnWidth(5, 130)
        self.gene_table.setColumnWidth(6, 220)
        self.gene_table.selectionModel().selectionChanged.connect(self._show_gene_detail)
        splitter.addWidget(self.gene_table)
        detail = QFrame()
        detail.setFrameShape(QFrame.Shape.StyledPanel)
        detail_layout = QVBoxLayout(detail)
        detail_top = QHBoxLayout()
        detail_top.addWidget(QLabel("详情与帮助"), 1)
        self.close_detail_button = QPushButton("收起详情")
        self.close_detail_button.clicked.connect(lambda: detail.setVisible(False))
        detail_top.addWidget(self.close_detail_button)
        detail_layout.addLayout(detail_top)
        self.gene_detail = QLabel("选择一行查看该特征的统计结果和各样本表达值。")
        self.gene_detail.setWordWrap(True)
        detail_layout.addWidget(self.gene_detail)
        self.expression_plot = QLabel("单基因表达分布将在此显示")
        self.expression_plot.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.expression_plot.setMinimumHeight(220)
        detail_layout.addWidget(self.expression_plot, 1)
        splitter.addWidget(detail)
        splitter.setSizes([850, 380])
        layout.addWidget(splitter, 1)
        table_actions = QHBoxLayout()
        self.show_detail_button = QPushButton("显示基因详情")
        self.show_detail_button.clicked.connect(lambda: detail.setVisible(True))
        table_actions.addWidget(self.show_detail_button)
        table_actions.addStretch(1)
        layout.addLayout(table_actions)
        self.tabs.addTab(page, "差异基因")

    def _build_information_tab(self) -> None:
        page = QWidget()
        layout = QVBoxLayout(page)
        self.analysis_info = QLabel("分析方法、输入信息及筛选设置将在分析完成后显示。")
        self.analysis_info.setWordWrap(True)
        layout.addWidget(self.analysis_info)
        layout.addWidget(QLabel("分析历史：每次正式重新运行都会保留原分析，不会静默覆盖。单纯调整筛选不会新增统计记录。"))
        row = QHBoxLayout()
        self.history_list = QListWidget()
        row.addWidget(self.history_list, 1)
        info = QVBoxLayout()
        self.history_detail = QLabel("选择一条历史记录查看其摘要。")
        self.history_detail.setWordWrap(True)
        info.addWidget(self.history_detail)
        self.open_history_button = QPushButton("打开所选历史记录")
        self.open_history_button.clicked.connect(self._open_selected_history)
        info.addWidget(self.open_history_button)
        info.addStretch(1)
        row.addLayout(info, 1)
        layout.addLayout(row, 1)
        self.history_list.currentRowChanged.connect(self._history_selected)
        self.tabs.addTab(page, "分析信息")

    @Slot(object)
    def set_matrix(self, matrix: ExpressionMatrix) -> None:
        self.matrix = matrix
        root = matrix.values.path.parent.parent
        self._report_path = root / "reports" / "analysis-report.html"
        self._result = None
        self._plan_data = None
        self._filter_is_applicable = False
        self.apply_filter_button.setEnabled(False)
        self.open_project_button.setEnabled(root.exists())
        self._refresh_history()

    @Slot(object)
    def set_result(self, result: object, plan: object | None = None) -> None:
        if self.matrix is None:
            return
        root = self.matrix.values.path.parent.parent
        self._report_path = root / "reports" / "analysis-report.html"
        self._result = dict(result) if isinstance(result, dict) else {}
        self._plan_data = plan.as_dict() if hasattr(plan, "as_dict") else None
        config = ResultFilterConfig.from_mapping(self._result.get("active_result_filter") or (self._plan_data or {}).get("result_filter"))
        self._set_filter_controls(config)
        self._filter_is_applicable = bool(self._plan_data and self._result.get("table"))
        self.apply_filter_button.setEnabled(self._filter_is_applicable)
        method = self._result.get("method", "")
        direction = ((self._plan_data or {}).get("contrast") or {})
        self.status_label.setText(
            f"分析完成 · 方法 {method} · 比较方向：{direction.get('test', 'test')} − {direction.get('control', 'control')}。"
            "log2FoldChange 为模型尺度差异，正值表示实验组较高。\n"
            f"项目目录：{root}\n{self._XLSX_HINT}"
        )
        sample_ids = (self._result.get("plot_matrix") or {}).get("sample_ids", [])
        summary = summarize_result(self._result.get("table", []), config)
        self.analysis_info.setText(
            f"输入矩阵：{self.matrix.values.path}\n"
            f"比较：{direction.get('test', 'test')} − {direction.get('control', 'control')}\n"
            f"方法：{method}；profile：{self._result.get('profile_id', '未记录')}\n"
            f"纳入样本：{len(sample_ids)}；检测特征：{summary['tested_count']}；"
            f"当前筛选显著：{summary['significant_magnitude_count']}（上调 {summary['up_count']}，下调 {summary['down_count']}）\n"
            f"当前阈值：{config.pvalue_mode} < {config.pvalue_threshold:g} 且 "
            f"|log2FC| ≥ {config.log2fc_threshold:g}；单位/尺度："
            f"{self.matrix.measurement_type.value} / {self.matrix.value_scale.value}\n"
            f"项目目录：{root}\n{self._XLSX_HINT}"
        )
        self.open_report_button.setEnabled(self._report_path.is_file())
        self.open_project_button.setEnabled(root.is_dir())
        self._load_charts(root)
        self._load_annotations()
        self._preview_filter()
        self._populate_gene_table()
        self._refresh_history()

    @Slot(object)
    def set_saved_project(self, snapshot: object) -> None:
        matrix = getattr(snapshot, "matrix", None)
        root = getattr(snapshot, "project_dir", None)
        if matrix is None or root is None:
            return
        self.matrix = matrix
        root = Path(root)
        self._report_path = getattr(snapshot, "report_path", None)
        result = getattr(snapshot, "saved_result", None)
        plan = getattr(snapshot, "saved_plan", None)
        if result is not None:
            self._result = dict(result)
            self._plan_data = dict(plan) if isinstance(plan, dict) else None
            saved_filter = (self._plan_data or {}).get("result_filter")
            filter_path = root / "analysis" / "result-filter.json"
            if filter_path.is_file() and not filter_path.is_symlink():
                try:
                    saved_filter = json.loads(filter_path.read_text(encoding="utf-8"))
                except (OSError, UnicodeError, json.JSONDecodeError):
                    pass
            self._set_filter_controls(ResultFilterConfig.from_mapping(saved_filter))
            self._filter_is_applicable = bool(getattr(snapshot, "result_integrity_valid", False) and self._plan_data and self._result.get("table"))
            self.apply_filter_button.setEnabled(self._filter_is_applicable)
            self._load_charts(root)
            self._load_annotations()
            self._preview_filter()
            self._populate_gene_table()
            active_config = self._current_filter()
            summary = summarize_result(self._result.get("table", []), active_config)
            direction = (self._plan_data or {}).get("contrast", {})
            sample_ids = (self._result.get("plot_matrix") or {}).get("sample_ids", [])
            self.analysis_info.setText(
                f"比较：{direction.get('test', 'test')} − {direction.get('control', 'control')}\n"
                f"方法：{self._result.get('method', '未知')}；profile：{self._result.get('profile_id', '未记录')}\n"
                f"纳入样本：{len(sample_ids)}；检测特征：{summary['tested_count']}；"
                f"当前筛选显著：{summary['significant_magnitude_count']}（上调 {summary['up_count']}，下调 {summary['down_count']}）\n"
                f"当前阈值：{active_config.pvalue_mode} < {active_config.pvalue_threshold:g} 且 "
                f"|log2FC| ≥ {active_config.log2fc_threshold:g}；"
                f"单位/尺度：{self.matrix.measurement_type.value} / {self.matrix.value_scale.value}\n"
                f"项目目录：{root}\n{self._XLSX_HINT}"
            )
            state = "已载入保存的统计结果；没有重新运行模型。"
        else:
            self._result = None
            self._plan_data = None
            state = "项目已载入，但没有可通过完整性核验的统计结果。"
        warnings = tuple(getattr(snapshot, "warnings", ()))
        if warnings:
            state += "\n核验提示：" + "；".join(warnings)
        self.status_label.setText(f"{state}\n项目目录：{root}")
        self.open_report_button.setEnabled(bool(getattr(snapshot, "report_valid", False) and self._report_path and self._report_path.is_file()))
        self.open_project_button.setEnabled(root.is_dir())
        self._refresh_history()

    def _current_filter(self) -> ResultFilterConfig:
        return ResultFilterConfig(
            pvalue_mode="fdr" if self.fdr_radio.isChecked() else "pvalue",
            pvalue_threshold=float(self.p_threshold.value()),
            log2fc_threshold=float(self.lfc_threshold.value()),
        )

    def _set_filter_controls(self, config: ResultFilterConfig) -> None:
        for widget in (self.fdr_radio, self.raw_p_radio, self.p_threshold, self.lfc_threshold):
            widget.blockSignals(True)
        self.fdr_radio.setChecked(config.pvalue_mode == "fdr")
        self.raw_p_radio.setChecked(config.pvalue_mode == "pvalue")
        self.p_threshold.setValue(config.pvalue_threshold)
        self.lfc_threshold.setValue(config.log2fc_threshold)
        for widget in (self.fdr_radio, self.raw_p_radio, self.p_threshold, self.lfc_threshold):
            widget.blockSignals(False)

    @Slot()
    def _preview_filter(self, *_args: object) -> None:
        if not self._result or not isinstance(self._result.get("table"), list):
            return
        config = self._current_filter()
        summary = summarize_result(self._result["table"], config)
        mode = "FDR" if config.pvalue_mode == "fdr" else "原始 P 值（未校正）"
        self.filter_summary.setText(
            f"检测 {summary['tested_count']} 个特征；{mode} < {config.pvalue_threshold:g} 且 "
            f"|log2FC| ≥ {config.log2fc_threshold:g}：{summary['significant_magnitude_count']} 个 "
            f"（上调 {summary['up_count']}，下调 {summary['down_count']}）。"
        )
        if self._result:
            current = self.analysis_info.text().split("\n")
            threshold_line = (
                f"当前阈值：{config.pvalue_mode} < {config.pvalue_threshold:g} 且 "
                f"|log2FC| ≥ {config.log2fc_threshold:g}；当前显著 {summary['significant_magnitude_count']} "
                f"（上调 {summary['up_count']}，下调 {summary['down_count']}）。"
            )
            current = [line for line in current if not line.startswith("当前阈值：")]
            current.append(threshold_line)
            self.analysis_info.setText("\n".join(current))
        self.filter_caveat.setText("原始 P 值未经过多重检验校正，使用时假阳性风险较高。" if config.pvalue_mode == "pvalue" else "筛选只影响展示，不改变统计模型或原始统计列。")
        self._populate_gene_table()

    def _visible_rows(self) -> list[dict[str, Any]]:
        if not self._result:
            return []
        config = self._current_filter()
        rows = [row for row in self._result.get("table", []) if isinstance(row, dict)]
        if self.hide_untested.isChecked():
            rows = [row for row in rows if row.get("tested") is True]
        if self.significant_only.isChecked():
            rows = [row for row in rows if passes_result_filter(row, config)]
        needle = self.gene_search.text().strip().casefold()
        if needle:
            rows = [row for row in rows if needle in str(row.get("feature_key") or row.get("original_id") or "").casefold() or needle in str(self._annotations.get(str(row.get("feature_key") or ""), row.get("gene_symbol") or "")).casefold()]
        return rows

    def _populate_gene_table(self, *_args: object) -> None:
        rows = self._result.get("table", []) if self._result else []
        self._gene_model.set_data(rows, self._annotations, self._current_filter())
        self._gene_proxy.config = self._current_filter()
        self._gene_proxy.query = self.gene_search.text().strip()
        self._gene_proxy.significant_only = self.significant_only.isChecked()
        self._gene_proxy.invalidateFilter()

    @Slot(str)
    def _set_gene_query(self, text: str) -> None:
        self._gene_proxy.query = text.strip()
        self._gene_proxy.invalidateFilter()

    @Slot(bool)
    def _set_significant_only(self, checked: bool) -> None:
        self._gene_proxy.significant_only = checked
        self._gene_proxy.invalidateFilter()

    @Slot(bool)
    def _set_hide_untested(self, checked: bool) -> None:
        self._gene_proxy.hide_untested = checked
        self._gene_proxy.invalidateFilter()

    @Slot()
    def _copy_gene_column(self, column: int) -> None:
        selection = self.gene_table.selectionModel().selectedRows()
        if not selection:
            return
        values = [str(self._gene_proxy.index(index.row(), column).data() or "") for index in selection]
        QGuiApplication.clipboard().setText("\n".join(value for value in values if value))
        self.filter_progress.setText(f"已复制 {len(values)} 条基因信息。")

    @staticmethod
    def _fmt(value: Any) -> str:
        if value is None or value == "":
            return "NA"
        try:
            number = float(value)
            return f"{number:.5g}"
        except (TypeError, ValueError):
            return str(value)

    @Slot(object, object)
    def _show_gene_detail(self, *_args: object) -> None:
        current = self.gene_table.currentIndex()
        if not current.isValid() or not self._result:
            return
        source_row = self._gene_proxy.mapToSource(current).row()
        row = self._gene_model.rows[source_row]
        feature_id = str(row.get("feature_key") or row.get("original_id") or "")
        plot = self._result.get("plot_matrix") or {}
        ids = plot.get("feature_ids", [])
        sample_ids = plot.get("sample_ids", [])
        matrix_values = plot.get("values", [])
        try:
            values = matrix_values[ids.index(feature_id)]
        except (ValueError, IndexError, TypeError):
            self.gene_detail.setText(f"{feature_id}\n没有可核验的逐样本变换后表达值。")
            return
        symbol = self._annotations.get(feature_id, row.get("gene_symbol") or "未注释")
        significant = passes_result_filter(row, self._current_filter())
        values_text = "；".join(f"{sample}: {float(value):.4g}" for sample, value in zip(sample_ids, values))
        measurement = self.matrix.measurement_type.value if self.matrix else "未知"
        scale = self.matrix.value_scale.value if self.matrix else "未知"
        self.gene_detail.setText(
            f"Gene ID：{feature_id}    Gene Symbol：{symbol}\n"
            f"log2FoldChange：{self._fmt(row.get('log2FoldChange'))}    原始 P：{self._fmt(row.get('pvalue'))}"
            f"    FDR：{self._fmt(row.get('padj'))}\n"
            f"当前筛选：{'显著' if significant else '未达到条件'}；条件为"
            f"{self._current_filter().pvalue_mode} < {self._current_filter().pvalue_threshold:g} 且 "
            f"|log2FC| ≥ {self._current_filter().log2fc_threshold:g}。\n"
            f"单位/输入尺度：{measurement} / {scale}；绘图值转换：{plot.get('transform', '未知')}。"
            f"逐样本表达值（分析保存尺度）：{values_text}。不另做统计检验。"
        )
        self._draw_expression(sample_ids, values)

    def _draw_expression(self, sample_ids: list[str], values: list[float]) -> None:
        from io import BytesIO
        from matplotlib.figure import Figure
        # Render in memory; no project/source file is modified and no Qt canvas backend is needed.
        figure = Figure(figsize=(7, 2.5), tight_layout=True)
        axis = figure.add_subplot(111)
        design = (self._plan_data or {}).get("design", {}).get("assignments", {})
        groups = [str(design.get(sample, "未分组")) for sample in sample_ids]
        palette = {"control": "#3977a8", "test": "#cb6948", "excluded": "#8b8b8b"}
        for index, (sample, value, group) in enumerate(zip(sample_ids, values, groups)):
            axis.scatter(index, float(value), color=palette.get(group, "#777777"), s=55)
            axis.text(index, float(value), sample, rotation=40, ha="right", va="bottom", fontsize=8)
        axis.set_ylabel("变换后表达值")
        axis.set_xticks([])
        axis.grid(axis="y", alpha=.2)
        buffer = BytesIO()
        figure.savefig(buffer, format="png", dpi=110)
        image = QPixmap()
        image.loadFromData(buffer.getvalue(), "PNG")
        self.expression_plot.setPixmap(image.scaled(self.expression_plot.size(), Qt.AspectRatioMode.KeepAspectRatio, Qt.TransformationMode.SmoothTransformation))

    def _load_charts(self, root: Path) -> None:
        self._chart_paths.clear()
        names = {"pca": "pca.png", "sample_correlation": "sample-correlation.png", "pvalue_histogram": "p-value-histogram.png", "volcano": "volcano.png", "top_heatmap": "top-heatmap.png"}
        for key, filename in names.items():
            path = root / "figures" / filename
            self._chart_paths[key] = path
            if path.is_file():
                pixmap = QPixmap(str(path))
                self._chart_labels[key].setPixmap(pixmap.scaled(self._chart_labels[key].size(), Qt.AspectRatioMode.KeepAspectRatio, Qt.TransformationMode.SmoothTransformation))
            else:
                self._chart_labels[key].setText("该项目没有此图表")

    def resizeEvent(self, event) -> None:
        super().resizeEvent(event)
        for key, path in self._chart_paths.items():
            if path.is_file():
                self._chart_labels[key].setPixmap(QPixmap(str(path)).scaled(self._chart_labels[key].size(), Qt.AspectRatioMode.KeepAspectRatio, Qt.TransformationMode.SmoothTransformation))

    def _open_chart(self, key: str) -> None:
        path = self._chart_paths.get(key)
        if path and path.is_file():
            QDesktopServices.openUrl(QUrl.fromLocalFile(str(path)))

    def _zoom_chart(self, key: str) -> None:
        path = self._chart_paths.get(key)
        if not path or not path.is_file():
            return
        dialog = QDialog(self)
        dialog.setWindowTitle(dict((item[0], item[1]) for item in self._CHARTS).get(key, "分析图表"))
        dialog.resize(max(720, int(self.width() * .9)), max(520, int(self.height() * .85)))
        layout = QVBoxLayout(dialog)
        image = QLabel()
        image.setAlignment(Qt.AlignmentFlag.AlignCenter)
        pixmap = QPixmap(str(path))
        image.setPixmap(pixmap.scaled(dialog.size(), Qt.AspectRatioMode.KeepAspectRatio, Qt.TransformationMode.SmoothTransformation))
        layout.addWidget(image, 1)
        close_button = QPushButton("关闭")
        close_button.clicked.connect(dialog.accept)
        layout.addWidget(close_button, 0, Qt.AlignmentFlag.AlignRight)
        dialog.exec()

    @Slot()
    def import_annotations(self) -> None:
        if self.matrix is None or not self._result or not self._plan_data:
            return
        path, _ = QFileDialog.getOpenFileName(self, "选择编号与 gene symbol 对照表", "", "表格 (*.tsv *.txt *.csv *.xlsx *.xlsm)")
        if not path:
            return
        try:
            source = Path(path)
            headers, rows = _read_annotation_table(source)
            feature_ids = {str(row.get("feature_key") or row.get("original_id") or "") for row in self._result.get("table", []) if isinstance(row, dict)}
            dialog = _AnnotationPreviewDialog(headers, rows, feature_ids, self)
            if dialog.exec() != QDialog.DialogCode.Accepted:
                return
            mapping = dialog.mapping
            if not mapping:
                raise ValueError("当前列选择下没有无歧义的 ID 与 Gene Symbol 映射。")
            self._start_annotation_update(mapping)
        except (OSError, UnicodeError, ValueError, csv.Error, KeyError) as exc:
            QMessageBox.warning(self, "无法读取注释表", str(exc))

    def _start_annotation_update(self, mapping: dict[str, str]) -> None:
        if self.matrix is None or self._plan_data is None:
            return
        if any(thread.isRunning() for thread, _worker, _context in self._annotation_threads):
            return
        context = JobContext(self.matrix.values.path.parent.parent)
        thread = QThread(self)
        worker = AnnotationWorker(self.matrix, self._plan_data, mapping, context)
        entry = (thread, worker, context)
        self._annotation_threads.append(entry)
        worker.moveToThread(thread)
        thread.started.connect(worker.run, Qt.ConnectionType.QueuedConnection)
        worker.progress.connect(self.filter_progress.setText, Qt.ConnectionType.QueuedConnection)
        worker.completed.connect(self._annotation_completed, Qt.ConnectionType.QueuedConnection)
        worker.failed.connect(self._annotation_failed, Qt.ConnectionType.QueuedConnection)
        worker.completed.connect(thread.quit, Qt.ConnectionType.QueuedConnection)
        worker.failed.connect(thread.quit, Qt.ConnectionType.QueuedConnection)
        thread.finished.connect(worker.deleteLater)
        thread.finished.connect(thread.deleteLater)
        thread.finished.connect(lambda current=entry: self._annotation_threads.remove(current))
        self.annotation_button.setEnabled(False)
        self.filter_progress.setText("正在更新注释、图表和报告；统计模型不会重跑……")
        thread.start()

    @Slot(object)
    def _annotation_completed(self, payload: object) -> None:
        data = dict(payload) if isinstance(payload, dict) else {}
        result = data.get("result")
        if isinstance(result, dict):
            self._result = result
        self._load_annotations()
        self._load_charts(self.matrix.values.path.parent.parent)
        self._populate_gene_table()
        self._preview_filter()
        self.open_report_button.setEnabled(bool(self._report_path and self._report_path.is_file()))
        self.annotation_button.setEnabled(True)
        self.filter_progress.setText(
            f"注释已应用：匹配 {data.get('matched_count', 0):,} 个编号；未匹配 {data.get('unmatched_count', 0):,} 个。"
            "相关图表、五列 XLSX 与 HTML 报告已更新，统计模型未重跑。"
        )

    @Slot(str)
    def _annotation_failed(self, message: str) -> None:
        self.annotation_button.setEnabled(True)
        self.filter_progress.setText(f"注释更新失败，现有结果保持不变：{message}")

    def _load_annotations(self) -> None:
        self._annotations = {}
        if self.matrix is None:
            return
        path = self.matrix.values.path.parent.parent / "analysis" / "gene-annotations.tsv"
        if path.is_file() and not path.is_symlink():
            try:
                with path.open(encoding="utf-8", newline="") as handle:
                    self._annotations = {row["feature_key"]: row["gene_symbol"] for row in csv.DictReader(handle, delimiter="\t") if row.get("feature_key") and row.get("gene_symbol")}
            except (OSError, UnicodeError, KeyError, csv.Error):
                self._annotations = {}

    @Slot()
    def export_filtered(self) -> None:
        if not self._result:
            return
        path, _ = QFileDialog.getSaveFileName(self, "导出当前筛选的差异基因", "filtered-de-results.tsv", "TSV (*.tsv)")
        if not path:
            return
        try:
            with Path(path).open("w", encoding="utf-8", newline="") as handle:
                writer = csv.writer(handle, delimiter="\t", lineterminator="\n")
                writer.writerow(("feature_key", "gene_symbol", "log2FoldChange", "pvalue", "padj"))
                for row in self._visible_rows():
                    key = str(row.get("feature_key") or row.get("original_id") or "")
                    writer.writerow((key, self._annotations.get(key, row.get("gene_symbol") or ""), row.get("log2FoldChange"), row.get("pvalue"), row.get("padj")))
            self.filter_progress.setText(f"已导出当前筛选结果：{path}")
        except OSError as exc:
            QMessageBox.warning(self, "导出失败", str(exc))

    def _refresh_history(self) -> None:
        self.history_list.clear()
        if self.matrix is None:
            return
        root = self.matrix.values.path.parent.parent / "analysis" / "history"
        if not root.is_dir():
            return
        for record in sorted(root.glob("*/history-manifest.json"), reverse=True):
            try:
                data = json.loads(record.read_text(encoding="utf-8"))
                label = f"{data.get('created_at', '未知时间')} · {data.get('contrast', '历史分析')} · {data.get('method', '')}"
                item = self.history_list.addItem(label)
                del item
                self.history_list.item(self.history_list.count() - 1).setData(Qt.ItemDataRole.UserRole, str(record.parent))
            except (OSError, UnicodeError, json.JSONDecodeError):
                continue

    @Slot(int)
    def _history_selected(self, index: int) -> None:
        item = self.history_list.item(index)
        if item is None:
            return
        folder = Path(item.data(Qt.ItemDataRole.UserRole))
        try:
            record = json.loads((folder / "history-manifest.json").read_text(encoding="utf-8"))
            contrast = record.get("contrast", "未知比较")
            summary = record.get("result_summary", {})
            groups = record.get("sample_assignments", {})
            control_n = sum(value == "control" for value in groups.values())
            test_n = sum(value == "test" for value in groups.values())
            details = (
                f"历史记录：{folder.name}\n生成时间：{record.get('created_at', '未知')}\n"
                f"来源：{record.get('source_path') or '原始来源路径未记录'}\n比较：{contrast}\n"
                f"方法：{record.get('method', '未知')}；样本数 control {control_n}、test {test_n}\n"
                f"检测特征：{summary.get('tested_count', '未知')}；筛选条件：{record.get('active_result_filter', {})}\n"
                f"注释状态：{'含后期注释文件' if record.get('annotation_status', {}).get('present') else '未补充后期注释'}\n"
                "历史报告和统计结果保存在该记录内，当前活动结果不受影响。"
            )
            self.history_detail.setText(details)
        except (OSError, UnicodeError, json.JSONDecodeError, AttributeError):
            self.history_detail.setText(f"历史记录：{folder.name}\n该记录缺少可读取的历史摘要；仍可尝试打开其文件夹。")

    @Slot()
    def _open_selected_history(self) -> None:
        item = self.history_list.currentItem()
        if item:
            folder = Path(item.data(Qt.ItemDataRole.UserRole))
            report = folder / "reports" / "analysis-report.html"
            QDesktopServices.openUrl(QUrl.fromLocalFile(str(report if report.is_file() else folder)))

    def apply_filter(self) -> None:
        if not self._filter_is_applicable or self.matrix is None or not self._plan_data or not self._result:
            return
        if any(thread.isRunning() for thread, _worker, _context in self._filter_threads):
            return
        context = JobContext(self.matrix.values.path.parent.parent)
        thread = QThread(self)
        worker = ResultFilterWorker(self.matrix, self._plan_data, self._current_filter(), context)
        entry = (thread, worker, context)
        self._filter_threads.append(entry)
        worker.moveToThread(thread)
        thread.started.connect(worker.run, Qt.ConnectionType.QueuedConnection)
        worker.progress.connect(self.filter_progress.setText, Qt.ConnectionType.QueuedConnection)
        worker.completed.connect(self._filter_completed, Qt.ConnectionType.QueuedConnection)
        worker.failed.connect(self._filter_failed, Qt.ConnectionType.QueuedConnection)
        worker.completed.connect(thread.quit, Qt.ConnectionType.QueuedConnection)
        worker.failed.connect(thread.quit, Qt.ConnectionType.QueuedConnection)
        thread.finished.connect(worker.deleteLater)
        thread.finished.connect(thread.deleteLater)
        thread.finished.connect(lambda current=entry: self._filter_threads.remove(current))
        self.apply_filter_button.setEnabled(False)
        self.filter_progress.setText("正在更新报告与图表；统计模型不会重跑……")
        thread.start()

    @Slot(object)
    def _filter_completed(self, value: object) -> None:
        result = dict(value) if isinstance(value, dict) else {}
        if self._result is not None:
            self._result["active_result_filter"] = result.get("result_filter")
            self._result["display_summary"] = result.get("summary")
        self.filter_progress.setText("筛选结果、图表与离线报告已更新；统计模型未重跑。")
        if self.matrix is not None:
            self._load_charts(self.matrix.values.path.parent.parent)
        self._preview_filter()
        self.open_report_button.setEnabled(bool(self._report_path and self._report_path.is_file()))
        self.apply_filter_button.setEnabled(self._filter_is_applicable)

    @Slot(str)
    def _filter_failed(self, message: str) -> None:
        self.filter_progress.setText(f"刷新失败，既有结果保持不变：{message}")
        self.apply_filter_button.setEnabled(self._filter_is_applicable)

    def active_threads(self) -> tuple[QThread, ...]:
        return tuple(thread for thread, _worker, _context in (*self._filter_threads, *self._annotation_threads))

    def cancel_refresh(self) -> None:
        for _thread, _worker, context in self._filter_threads:
            context.cancel_event.set()
        for _thread, _worker, context in self._annotation_threads:
            context.cancel_event.set()

    @Slot()
    def open_report(self) -> None:
        if self._report_path and self._report_path.exists():
            QDesktopServices.openUrl(QUrl.fromLocalFile(str(self._report_path)))

    @Slot()
    def open_project(self) -> None:
        if self.matrix is not None:
            root = self.matrix.values.path.parent.parent
            if root.exists():
                QDesktopServices.openUrl(QUrl.fromLocalFile(str(root)))
