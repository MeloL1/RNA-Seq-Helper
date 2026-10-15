from __future__ import annotations

import csv
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Callable, Iterable

from openpyxl import load_workbook
import xlrd


@dataclass(frozen=True)
class RawTable:
    path: Path
    sheet_name: str | None
    rows: tuple[tuple[str, ...], ...]
    row_count: int
    column_count: int
    encoding: str | None = None
    delimiter: str | None = None
    formula_cells: tuple[str, ...] = ()
    is_complete: bool = True


class TableReadCancelled(RuntimeError):
    """用户在后台读取表格期间请求停止。"""


class TableReadLimitError(ValueError):
    """预览记录超过资源上限，不能在安全范围内解析。"""


class TableReader:
    TEXT_SUFFIXES = {".csv", ".tsv", ".txt", ".tab"}
    EXCEL_SUFFIXES = {".xlsx", ".xlsm", ".xls"}
    MAX_PREVIEW_PHYSICAL_LINE_CHARACTERS = 1_048_576
    MAX_PREVIEW_RETAINED_CHARACTERS = 2_000_000

    def list_sheets(self, path: Path) -> tuple[str | None, ...]:
        if path.suffix.lower() in self.TEXT_SUFFIXES:
            return (None,)
        if path.suffix.lower() in self.EXCEL_SUFFIXES:
            if path.suffix.lower() == ".xls":
                if self._looks_like_delimited_text(path):
                    # 部分测序服务商把制表符文本命名为 .xls；它不是
                    # BIFF Excel，不能交给 xlrd，但可以按文本表格读取。
                    return (None,)
                workbook = xlrd.open_workbook(str(path), on_demand=True)
                try:
                    return tuple(workbook.sheet_names())
                finally:
                    workbook.release_resources()
            workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
            try:
                return tuple(workbook.sheetnames)
            finally:
                workbook.close()
        raise ValueError(f"不支持的表格格式：{path.suffix or '无扩展名'}")

    def read(
        self,
        path: Path,
        sheet_name: str | None = None,
        max_rows: int | None = None,
        max_columns: int | None = None,
        *,
        progress: Callable[[int], None] | None = None,
        cancelled: Callable[[], bool] | None = None,
    ) -> RawTable:
        suffix = path.suffix.lower()
        if suffix in self.TEXT_SUFFIXES:
            return self._read_text(path, max_rows, max_columns, progress, cancelled)
        if suffix in self.EXCEL_SUFFIXES:
            if suffix == ".xls":
                if self._looks_like_delimited_text(path):
                    return self._read_text(path, max_rows, max_columns, progress, cancelled)
                return self._read_xls(path, sheet_name, max_rows, max_columns, progress, cancelled)
            return self._read_excel(path, sheet_name, max_rows, max_columns, progress, cancelled)
        raise ValueError(f"不支持的表格格式：{path.suffix or '无扩展名'}")

    def _read_text(
        self,
        path: Path,
        max_rows: int | None,
        max_columns: int | None,
        progress: Callable[[int], None] | None = None,
        cancelled: Callable[[], bool] | None = None,
    ) -> RawTable:
        # 预览只读取有限前缀；不能为显示前 40 行而先把整个大表读进内存。
        with path.open("rb") as handle:
            prefix = handle.read(131072)
        encoding = self._detect_encoding(prefix)
        sample = prefix.decode(encoding, errors="ignore")[:8192]
        delimiter = self._detect_delimiter(sample, path)
        rows = []
        row_limit_reached = False
        columns_truncated = False
        retained_characters = 0
        with path.open("r", encoding=encoding, errors="strict", newline="") as handle:
            preview_bounded = max_rows is not None or max_columns is not None
            source = self._bounded_preview_lines(handle) if preview_bounded else handle
            reader = csv.reader(source, delimiter=delimiter)
            for index, row in enumerate(self._iter_csv_rows(reader, preview_bounded)):
                if cancelled and cancelled():
                    raise TableReadCancelled("已取消读取表格")
                if max_rows is not None and index >= max_rows:
                    row_limit_reached = True
                    break
                if max_columns is not None and len(row) > max_columns:
                    columns_truncated = True
                selected = row if max_columns is None else row[:max_columns]
                if preview_bounded:
                    retained_characters += sum(len(value) for value in selected)
                    if retained_characters > self.MAX_PREVIEW_RETAINED_CHARACTERS:
                        raise TableReadLimitError(
                            "预览内容超过安全上限。请先删去超长说明列或把表格另存为仅含本次基因和样本的 CSV/TSV，再导入。"
                        )
                rows.append(tuple(self._display_value(value) for value in selected))
                if progress and (index + 1) % 1000 == 0:
                    progress(index + 1)
        column_count = max((len(row) for row in rows), default=0)
        return RawTable(
            path, None, tuple(rows), len(rows), column_count, encoding, delimiter,
            is_complete=not row_limit_reached and not columns_truncated,
        )

    def _bounded_preview_lines(self, handle):
        """避免单个异常宽物理行在 csv.reader 中形成无界列表。"""
        while True:
            line = handle.readline(self.MAX_PREVIEW_PHYSICAL_LINE_CHARACTERS + 1)
            if not line:
                return
            if len(line) > self.MAX_PREVIEW_PHYSICAL_LINE_CHARACTERS:
                raise TableReadLimitError(
                    "表格有一行超过安全预览上限（约 1 Mi 字符）。请移除无关列或另存为精简的 CSV/TSV 后重试。"
                )
            yield line

    @staticmethod
    def _iter_csv_rows(reader, preview_bounded: bool):
        try:
            yield from reader
        except csv.Error as exc:
            if preview_bounded and "field larger than field limit" in str(exc).casefold():
                raise TableReadLimitError(
                    "表格预览中有单个格子的内容超过安全预览上限。请删去超长说明列或改用精简的 CSV/TSV。"
                ) from exc
            raise

    def _read_excel(
        self,
        path: Path,
        sheet_name: str | None,
        max_rows: int | None,
        max_columns: int | None,
        progress: Callable[[int], None] | None = None,
        cancelled: Callable[[], bool] | None = None,
    ) -> RawTable:
        sheets = self.list_sheets(path)
        selected = sheet_name or (sheets[0] if sheets else None)
        if selected not in sheets:
            raise ValueError(f"找不到工作表：{selected}")
        workbook_values = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        workbook_formulas = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            values_sheet = workbook_values[selected]
            formula_sheet = workbook_formulas[selected]
            source_column_count = values_sheet.max_column or 0
            read_max_column = (
                max_columns
                if max_columns is not None and source_column_count > max_columns
                else None
            )
            rows: list[tuple[str, ...]] = []
            formulas: list[str] = []
            row_limit_reached = False
            columns_truncated = max_columns is not None and source_column_count > max_columns
            row_iterator = (
                values_sheet.iter_rows(max_col=read_max_column)
                if read_max_column is not None
                else values_sheet.iter_rows()
            )
            for row_index, row in enumerate(row_iterator):
                if cancelled and cancelled():
                    raise TableReadCancelled("已取消读取工作簿")
                if max_rows is not None and row_index >= max_rows:
                    row_limit_reached = True
                    break
                if max_columns is not None and len(row) > max_columns:
                    columns_truncated = True
                selected_cells = row[:max_columns] if max_columns is not None else row
                values = tuple(self._display_value(cell.value) for cell in selected_cells)
                rows.append(values)
                if row_index < 1000:
                    formula_row = (
                        next(formula_sheet.iter_rows(
                            min_row=row_index + 1,
                            max_row=row_index + 1,
                            max_col=read_max_column,
                        ))
                        if read_max_column is not None
                        else formula_sheet[row_index + 1]
                    )
                    for cell in formula_row[:max_columns] if max_columns is not None else formula_row:
                        if isinstance(cell.value, str) and cell.value.startswith("="):
                            formulas.append(cell.coordinate)
                if progress and (row_index + 1) % 1000 == 0:
                    progress(row_index + 1)
            # column_count tracks the materialized preview width. Keep the source
            # width only for deciding whether this bounded read is incomplete;
            # PreparationService compares preview width to the later full read.
            column_count = max((len(row) for row in rows), default=0)
            return RawTable(
                path, selected, tuple(rows), len(rows), column_count,
                formula_cells=tuple(formulas),
                is_complete=not row_limit_reached and not columns_truncated,
            )
        finally:
            workbook_values.close()
            workbook_formulas.close()

    def _read_xls(
        self,
        path: Path,
        sheet_name: str | None,
        max_rows: int | None,
        max_columns: int | None,
        progress: Callable[[int], None] | None = None,
        cancelled: Callable[[], bool] | None = None,
    ) -> RawTable:
        sheets = self.list_sheets(path)
        selected = sheet_name or (sheets[0] if sheets else None)
        if selected not in sheets:
            raise ValueError(f"找不到工作表：{selected}")
        workbook = xlrd.open_workbook(str(path), on_demand=True)
        try:
            sheet = workbook.sheet_by_name(selected)
            row_limit = min(sheet.nrows, max_rows) if max_rows is not None else sheet.nrows
            column_limit = min(sheet.ncols, max_columns) if max_columns is not None else sheet.ncols
            rows: list[tuple[str, ...]] = []
            for row_index in range(row_limit):
                if cancelled and cancelled():
                    raise TableReadCancelled("已取消读取 XLS 工作簿")
                values = tuple(
                    self._xls_display_value(sheet.cell(row_index, column), workbook.datemode)
                    for column in range(column_limit)
                )
                rows.append(values)
                if progress and (row_index + 1) % 1000 == 0:
                    progress(row_index + 1)
            return RawTable(
                path,
                selected,
                tuple(rows),
                len(rows),
                column_limit,
                is_complete=row_limit >= sheet.nrows and column_limit >= sheet.ncols,
            )
        finally:
            workbook.release_resources()

    @staticmethod
    def _xls_display_value(cell: xlrd.sheet.Cell, datemode: int) -> str:
        if cell.ctype == xlrd.XL_CELL_EMPTY:
            return ""
        if cell.ctype == xlrd.XL_CELL_DATE:
            return xlrd.xldate_as_datetime(cell.value, datemode).isoformat()
        return TableReader._display_value(cell.value)

    @staticmethod
    def _detect_encoding(raw: bytes) -> str:
        for marker, encoding in ((b"\xef\xbb\xbf", "utf-8-sig"), (b"\xff\xfe", "utf-16"), (b"\xfe\xff", "utf-16")):
            if raw.startswith(marker):
                return encoding
        for encoding in ("utf-8", "gb18030", "utf-16"):
            # 前缀可能恰好在多字节字符中间结束；最多丢弃末尾 4 字节再判。
            for trim in range(5):
                sample = raw[:-trim] if trim else raw
                try:
                    sample.decode(encoding)
                    return encoding
                except UnicodeDecodeError:
                    continue
        raise ValueError("无法确认文本编码，请导出为 UTF-8 或 GB18030 后重试")

    @classmethod
    def _looks_like_delimited_text(cls, path: Path) -> bool:
        """识别扩展名为 .xls、实际内容为文本表格的服务商结果。"""

        try:
            with path.open("rb") as handle:
                sample_raw = handle.read(65536)
            if not sample_raw or b"\x00" in sample_raw:
                return False
            encoding = cls._detect_encoding(sample_raw)
            sample = sample_raw.decode(encoding, errors="strict")
            first_line = next((line for line in sample.splitlines() if line.strip()), "")
            if "\t" not in first_line and "," not in first_line and ";" not in first_line and "|" not in first_line:
                return False
            return "\n" in sample or "\r" in sample
        except (OSError, UnicodeDecodeError, ValueError):
            return False

    @staticmethod
    def _detect_delimiter(sample: str, path: Path) -> str:
        if path.suffix.lower() in {".tsv", ".tab"}:
            return "\t"
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
            return dialect.delimiter
        except csv.Error:
            return "\t" if "\t" in sample else ","

    @staticmethod
    def _display_value(value: object) -> str:
        if value is None:
            return ""
        if isinstance(value, (datetime, date)):
            return value.isoformat()
        return str(value)
