from __future__ import annotations

import csv
import hashlib
import html.parser
import json
import math
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from rnaseq_helper.domain.models import (
    ExpressionMatrix,
    Issue,
    IssueSeverity,
    MatrixStoreRef,
    MeasurementType,
    ObservationUnit,
    ValidationSummary,
    ValueScale,
    sha256_file,
)
from rnaseq_helper.services.project_lock import ProjectBusyError, project_directory_lock


class ProjectLoadError(ValueError):
    """保存项目格式不完整、路径越界或关键文件校验失败。"""


class ProjectLoadCancelled(RuntimeError):
    """用户取消项目来源文件校验。"""


_RESULT_EXPORT_COLUMNS = (
    "feature_key",
    "gene_symbol",
    "log2FoldChange",
    "pvalue",
    "padj",
)
_FIGURE_FILES = {
    "pca": "pca.png",
    "sample_correlation": "sample-correlation.png",
    "pvalue_histogram": "p-value-histogram.png",
    "volcano": "volcano.png",
    "top_heatmap": "top-heatmap.png",
}


class _ReportHTMLValidator(html.parser.HTMLParser):
    """对离线报告执行轻量结构检查，不运行脚本或加载外部资源。"""

    _VOID_TAGS = {
        "area", "base", "br", "col", "embed", "hr", "img", "input", "link",
        "meta", "param", "source", "track", "wbr",
    }

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.tags: set[str] = set()
        self.stack: list[str] = []
        self.has_body_content = False
        self.invalid_nesting = False

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        del attrs
        tag = tag.casefold()
        self.tags.add(tag)
        if tag not in self._VOID_TAGS:
            self.stack.append(tag)

    def handle_endtag(self, tag: str) -> None:
        tag = tag.casefold()
        if not self.stack or self.stack[-1] != tag:
            self.invalid_nesting = True
            return
        self.stack.pop()

    def handle_data(self, data: str) -> None:
        if data.strip() and "body" in self.stack:
            self.has_body_content = True

    def is_valid_report(self) -> bool:
        return (
            {"html", "head", "body"}.issubset(self.tags)
            and not self.stack
            and not self.invalid_nesting
            and self.has_body_content
        )


@dataclass(frozen=True)
class ProjectSnapshot:
    """项目快照；不整理或分析，加载时可能先恢复中断发布事务。"""

    project_dir: Path
    matrix: ExpressionMatrix
    saved_plan: dict[str, Any] | None
    saved_result: dict[str, Any] | None
    warnings: tuple[str, ...] = ()
    report_path: Path | None = None
    report_valid: bool = False
    result_integrity_valid: bool = False


class ProjectStore:
    """读取已保存矩阵和产物并校验；发现中断发布时先恢复旧结果。"""

    @classmethod
    def load(
        cls,
        project_dir: Path,
        *,
        progress=None,
        cancelled=None,
    ) -> ProjectSnapshot:
        root = Path(project_dir).expanduser().resolve(strict=True)
        if not root.is_dir():
            raise ProjectLoadError("请选择一个项目文件夹。")
        manifest_candidate = root / "project.json"
        if manifest_candidate.is_symlink() or not manifest_candidate.is_file():
            raise ProjectLoadError("这个文件夹不是可识别的 RNA-Seq Helper 项目；没有修改该文件夹。")
        cls._json(manifest_candidate, "项目清单 project.json")
        try:
            with project_directory_lock(root):
                return cls._load_locked(root, progress=progress, cancelled=cancelled)
        except ProjectBusyError as exc:
            raise ProjectLoadError(str(exc)) from exc

    @classmethod
    def _load_locked(cls, root: Path, *, progress=None, cancelled=None) -> ProjectSnapshot:
        cls._raise_if_cancelled(cancelled)
        try:
            from rnaseq_helper.analysis.service import AnalysisService

            transaction_warnings = list(AnalysisService.recover_interrupted_outputs(root))
        except ValueError as exc:
            raise ProjectLoadError(str(exc)) from exc
        cls._raise_if_cancelled(cancelled)
        manifest_path = cls._member(root, "project.json")
        matrix_path = cls._member(root, "prepared/expression_matrix.tsv")
        feature_path = cls._member(root, "prepared/feature_metadata.tsv")
        sample_path = cls._member(root, "prepared/sample_metadata.tsv")

        manifest = cls._json(manifest_path, "项目清单 project.json")
        recipe_candidate = root / "preparation-recipe.json"
        recipe = (
            cls._json(cls._member(root, "preparation-recipe.json"), "整理配方 preparation-recipe.json")
            if recipe_candidate.is_file() else {}
        )
        warnings: list[str] = transaction_warnings
        source_path_value = manifest.get("source_path")
        source_path = None
        if isinstance(source_path_value, str) and source_path_value.strip():
            if source_path_value.startswith("["):
                warnings.append("项目清单已隐藏原始来源路径；只能核验保存的整理矩阵和结果，无法检查原始文件是否仍在。")
            elif manifest.get("source_path_relative") is True:
                try:
                    source_path = cls._member(root, source_path_value)
                    source_exists = True
                except ProjectLoadError:
                    source_path = None
                    source_exists = False
                if not source_exists:
                    warnings.append("项目内保存的来源文件无法访问；整理矩阵和既有结果仍可只读查看，但无法从来源文件复核。")
                else:
                    recorded_source_hash = manifest.get("source_sha256")
                    if isinstance(recorded_source_hash, str) and recorded_source_hash:
                        try:
                            from rnaseq_helper.preparation.service import (
                                PreparationCancelled,
                                PreparationService,
                            )

                            source_matches = PreparationService.fingerprint_source(
                                source_path,
                                progress=progress,
                                cancelled=cancelled,
                            ) == recorded_source_hash
                        except PreparationCancelled as exc:
                            raise ProjectLoadCancelled("已取消来源文件校验；未继续打开项目。") from exc
                        except (OSError, ValueError):
                            source_matches = False
                        if not source_matches:
                            warnings.append("项目内来源文件与导入时记录的校验值不一致；整理矩阵和既有结果仍可只读查看，但不能把该文件当作原件复核。")
                    else:
                        warnings.append("项目未记录来源文件校验值；整理矩阵和既有结果仍可只读查看，但无法确认来源内容是否变化。")
            else:
                try:
                    source_path = Path(source_path_value).expanduser()
                    source_exists = source_path.is_absolute() and source_path.is_file()
                except (OSError, RuntimeError, ValueError):
                    source_path = None
                    source_exists = False
                if not source_exists:
                    warnings.append("原始来源文件当前无法访问；保存的整理矩阵和既有结果仍可只读查看，但无法从原始文件复核。")
                else:
                    recorded_source_hash = manifest.get("source_sha256")
                    if isinstance(recorded_source_hash, str) and recorded_source_hash:
                        try:
                            from rnaseq_helper.preparation.service import (
                                PreparationCancelled,
                                PreparationService,
                            )

                            source_matches = PreparationService.fingerprint_source(
                                source_path,
                                progress=progress,
                                cancelled=cancelled,
                            ) == recorded_source_hash
                        except PreparationCancelled as exc:
                            raise ProjectLoadCancelled("已取消来源文件校验；未继续打开项目。") from exc
                        except (OSError, ValueError):
                            source_matches = False
                        if not source_matches:
                            warnings.append("原始来源文件内容与导入时记录不一致；保存的整理矩阵和既有结果仍可只读查看，但不能把当前来源当作原件复核。")
                    else:
                        warnings.append("项目未记录原始来源文件校验值；保存的整理矩阵和既有结果仍可只读查看，但无法确认来源内容是否变化。")
        if not recipe:
            warnings.append("项目没有保存整理配方；数值类型与尺度按 UNKNOWN 处理，仅允许只读查看。")
        try:
            measurement_type = MeasurementType(str(recipe.get("measurement_type", "UNKNOWN")))
            value_scale = ValueScale(str(recipe.get("value_scale", "UNKNOWN")))
        except (KeyError, ValueError) as exc:
            raise ProjectLoadError("项目没有有效保存数值类型和尺度；不能猜测这些信息。") from exc
        confirmations = recipe.get("fact_confirmations", ())
        confirmed_values = {
            str(item.get("fact_key", "")): str(item.get("selected_value", ""))
            for item in confirmations
            if isinstance(item, dict)
            and str(item.get("selected_value", "")) not in {"", "UNKNOWN"}
            and str(item.get("status", "")) != "not_confirmed"
        } if isinstance(confirmations, list) else {}
        if measurement_type != MeasurementType.UNKNOWN and confirmed_values.get("measurement_type") != measurement_type.value:
            measurement_type = MeasurementType.UNKNOWN
            warnings.append("旧配方中的数值类型没有对应的确认记录；保留旧结果供查看，但新分析将按 UNKNOWN 阻断。")
        if value_scale != ValueScale.UNKNOWN and confirmed_values.get("value_scale") != value_scale.value:
            value_scale = ValueScale.UNKNOWN
            warnings.append("旧配方中的数值尺度没有对应的确认记录；保留旧结果供查看，但新分析将按 UNKNOWN 阻断。")

        matrix_digest = cls._sha256_file(matrix_path, progress=progress, cancelled=cancelled)
        cls._raise_if_cancelled(cancelled)
        manifest_digest = manifest.get("matrix_content_hash")
        if manifest_digest:
            if str(manifest_digest) != matrix_digest:
                raise ProjectLoadError("整理矩阵与项目清单中的校验值不一致；为避免使用被改动的数据，已停止打开。")
        else:
            warnings.append("旧项目没有保存矩阵校验值；已做结构检查，但完整性校验有限。")

        feature_ids, sample_ids, numeric_count = cls._read_expression_matrix(matrix_path, cancelled=cancelled)
        cls._check_feature_metadata(feature_path, feature_ids, cancelled=cancelled)
        cls._check_sample_metadata(sample_path, sample_ids, cancelled=cancelled)
        if not feature_ids or not sample_ids:
            raise ProjectLoadError("项目矩阵为空，无法恢复。")

        validation = cls._validation_summary(
            manifest.get("validation"), len(feature_ids), len(sample_ids), numeric_count, warnings,
        )
        if not validation.passed:
            raise ProjectLoadError("项目的整理校验没有通过；不能将其作为已准备好的分析项目恢复。")
        source_asset_id = str(manifest.get("source_asset_id", "")) or f"legacy-asset-{matrix_digest[:16]}"
        matrix_id = str(manifest.get("matrix_id", "")) or f"legacy-matrix-{matrix_digest[:16]}"
        feature_id_type = str(manifest.get("feature_id_type", "")) or "UNKNOWN"
        source_asset = manifest.get("source_asset")
        source_unit_hint = source_asset.get("observation_unit_hint") if isinstance(source_asset, dict) else None
        unit_values = [str(value).strip().upper() for value in (
            manifest.get("observation_unit"), source_unit_hint,
        ) if value not in (None, "")]
        try:
            recorded_units = [ObservationUnit(value) for value in unit_values]
        except ValueError as exc:
            raise ProjectLoadError("项目记录了不支持的样本观察单位；不能安全恢复用于分析。") from exc
        if ObservationUnit.CELL in recorded_units:
            # 任一持久化来源明确标记细胞级，就不能被另一个被编辑/旧字段降级为 bulk。
            observation_unit = ObservationUnit.CELL
        elif ObservationUnit.UNKNOWN in recorded_units or len(set(recorded_units)) > 1:
            observation_unit = ObservationUnit.UNKNOWN
        else:
            observation_unit = recorded_units[0] if recorded_units else ObservationUnit.UNKNOWN
        for field, value in (
            ("source_asset_id", manifest.get("source_asset_id")),
            ("matrix_id", manifest.get("matrix_id")),
            ("feature_id_type", manifest.get("feature_id_type")),
        ):
            if not value:
                warnings.append(f"旧项目缺少 {field}；仅生成临时恢复标识或 UNKNOWN，不会改写原项目。")

        recipe_id = hashlib.sha256(
            json.dumps(recipe, ensure_ascii=False, sort_keys=True).encode("utf-8")
        ).hexdigest()[:16]
        source_modality = manifest.get("source_modality")
        modality_inference_reason = ""
        if not source_modality:
            source_modality, modality_inference_reason = cls._infer_legacy_source_modality(
                manifest, sample_path,
            )
            if source_modality:
                warnings.append(
                    "旧项目没有保存来源类别；根据其保存的 GEO 来源元数据，"
                    f"现按“{source_modality}”门禁恢复（{modality_inference_reason}）。"
                    "既有结果仅作为历史记录；重新分析前须按来源说明重新确认数值类型和尺度。"
                )
        matrix = ExpressionMatrix(
            matrix_id=matrix_id,
            values=MatrixStoreRef(
                matrix_path, "dense-tsv", "float64",
                (len(feature_ids), len(sample_ids)), matrix_digest,
            ),
            feature_metadata=MatrixStoreRef(
                feature_path, "table-tsv", "string",
                (len(feature_ids), cls._column_count(feature_path)), sha256_file(feature_path),
            ),
            sample_metadata=MatrixStoreRef(
                sample_path, "table-tsv", "string",
                (len(sample_ids), cls._column_count(sample_path)), sha256_file(sample_path),
            ),
            observation_unit=observation_unit,
            measurement_type=measurement_type,
            value_scale=value_scale,
            input_transform=str(recipe.get("input_transform", "")),
            feature_id_type=feature_id_type,
            species_taxid=str(recipe["species_taxid"]) if recipe.get("species_taxid") else None,
            source_asset_ids=(source_asset_id,),
            recipe_id=recipe_id,
            validation=validation,
            content_hash=matrix_digest,
            measurement_provenance=cls._measurement_provenance(
                recipe.get("fact_confirmations", ()),
            ),
            platform_id=str(manifest["platform_id"]) if manifest.get("platform_id") else None,
            source_modality=str(source_modality) if source_modality else None,
            source_path=source_path,
        )

        saved_plan = cls._load_plan(root, matrix_digest, feature_path, sample_ids, warnings)
        cls._raise_if_cancelled(cancelled)
        saved_result, report_path, report_valid, result_integrity_valid = cls._load_result(
            root, feature_ids, matrix_digest, feature_path, sample_path, saved_plan, warnings,
        )
        if saved_plan is not None and not cls._has_confirmed_replicate_structure(
            saved_plan.get("design"),
        ):
            warnings.append(
                "旧分析计划没有保存样本独立性/配对确认；既有结果仍可只读查看，但旧 HTML 报告不作为已核验报告打开，重新分析前须确认设计。"
            )
            report_valid = False
        cls._raise_if_cancelled(cancelled)
        return ProjectSnapshot(
            root, matrix, saved_plan, saved_result, tuple(warnings), report_path, report_valid,
            result_integrity_valid,
        )

    @staticmethod
    def _has_confirmed_replicate_structure(design: Any) -> bool:
        if not isinstance(design, dict):
            return False
        structure = design.get("replicate_structure")
        expected_type = "paired" if structure == "paired" else "independent_two_group"
        return (
            structure in {"independent", "paired"}
            and design.get("replicate_structure_confirmed") is True
            and design.get("type") == expected_type
        )

    @classmethod
    def _infer_legacy_source_modality(
        cls,
        manifest: dict[str, Any],
        sample_metadata_path: Path,
    ) -> tuple[str | None, str]:
        """Recover only high-confidence array modality clues omitted by legacy manifests."""
        source_asset = manifest.get("source_asset")
        if not isinstance(source_asset, dict) or str(source_asset.get("kind", "")).upper() != "GEO_SERIES_MATRIX":
            return None, ""
        series_types = manifest.get("series_types", ())
        if isinstance(series_types, str):
            series_types = (series_types,)
        if isinstance(series_types, (list, tuple)) and any(
            str(value).strip().casefold() == "expression profiling by array"
            for value in series_types
        ):
            return "MICROARRAY", "保存的 GEO Series_type 标注为 Expression profiling by array"

        try:
            with sample_metadata_path.open("r", encoding="utf-8-sig", newline="") as handle:
                reader = csv.DictReader(handle, delimiter="\t")
                processing_values: list[str] = []
                supplementary_values: list[str] = []
                for row in reader:
                    processing_values.extend(
                        str(value or "") for key, value in row.items()
                        if key and key.strip().casefold() in {"data_processing", "series_type"}
                    )
                    supplementary_values.extend(
                        str(value or "") for key, value in row.items()
                        if key and key.strip().casefold() in {"supplementary_file", "supplementary_files"}
                    )
        except (OSError, UnicodeError, csv.Error):
            return None, ""
        has_affymetrix_processing = any("affymetrix" in value.casefold() for value in processing_values)
        has_cel_source = any(".cel" in value.casefold() for value in supplementary_values)
        if has_affymetrix_processing and has_cel_source:
            return (
                "MICROARRAY",
                "保存的 GEO 样本注释同时含 Affymetrix 处理说明和 CEL 来源文件",
            )
        return None, ""

    @staticmethod
    def _member(root: Path, relative: str) -> Path:
        path = root / relative
        if path.is_symlink():
            raise ProjectLoadError(f"项目关键文件不能是符号链接：{relative}")
        try:
            resolved = path.resolve(strict=True)
        except OSError as exc:
            raise ProjectLoadError(f"项目缺少必需文件：{relative}") from exc
        if not resolved.is_relative_to(root) or not resolved.is_file():
            raise ProjectLoadError(f"项目文件路径无效或指向项目文件夹之外：{relative}")
        return resolved

    @staticmethod
    def _json(path: Path, label: str) -> dict[str, Any]:
        try:
            value = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, UnicodeError, json.JSONDecodeError) as exc:
            raise ProjectLoadError(f"{label}损坏或无法读取。") from exc
        if not isinstance(value, dict):
            raise ProjectLoadError(f"{label}格式无效。")
        return value

    @staticmethod
    def _read_expression_matrix(path: Path, *, cancelled=None) -> tuple[list[str], list[str], int]:
        features: list[str] = []
        numeric_count = 0
        try:
            with path.open(encoding="utf-8", newline="") as handle:
                reader = csv.reader(handle, delimiter="\t")
                header = next(reader, None)
                if not header or len(header) < 3 or header[0].strip() != "feature_key":
                    raise ProjectLoadError("项目表达矩阵表头无效；应以 feature_key 开头且至少包含两个样本。")
                samples = header[1:]
                if any(not value or value != value.strip() for value in samples) or len(set(samples)) != len(samples):
                    raise ProjectLoadError("项目表达矩阵的样本编号为空或重复。")
                for line_number, row in enumerate(reader, start=2):
                    if line_number % 1024 == 0:
                        ProjectStore._raise_if_cancelled(cancelled)
                    if len(row) != len(header):
                        raise ProjectLoadError(f"项目表达矩阵第 {line_number} 行的列数不一致。")
                    feature_id = row[0]
                    if not feature_id or feature_id != feature_id.strip():
                        raise ProjectLoadError(f"项目表达矩阵第 {line_number} 行缺少 feature_key。")
                    features.append(feature_id)
                    for value in row[1:]:
                        number = float(value)
                        if not math.isfinite(number):
                            raise ProjectLoadError(f"项目表达矩阵第 {line_number} 行含非有限数值。")
                        numeric_count += 1
        except (OSError, UnicodeError, csv.Error, ValueError, StopIteration) as exc:
            if isinstance(exc, ProjectLoadError):
                raise
            raise ProjectLoadError("项目表达矩阵损坏或包含无效数值。") from exc
        if len(features) != len(set(features)):
            raise ProjectLoadError("项目表达矩阵存在重复 feature_key，不能安全恢复。")
        return features, samples, numeric_count

    @staticmethod
    def _check_feature_metadata(path: Path, feature_ids: list[str], *, cancelled=None) -> None:
        try:
            with path.open(encoding="utf-8", newline="") as handle:
                reader = csv.reader(handle, delimiter="\t")
                header = next(reader, None)
                if not header or "feature_key" not in header:
                    raise ProjectLoadError("项目特征注释表缺少 feature_key 列。")
                key_index = header.index("feature_key")
                keys = []
                for line_number, row in enumerate(reader, start=2):
                    if line_number % 1024 == 0:
                        ProjectStore._raise_if_cancelled(cancelled)
                    if len(row) > key_index:
                        keys.append(row[key_index])
        except (OSError, UnicodeError, csv.Error, StopIteration) as exc:
            raise ProjectLoadError("项目特征注释表损坏或无法读取。") from exc
        if keys != feature_ids:
            raise ProjectLoadError("特征注释编号或顺序与表达矩阵不一致。")

    @staticmethod
    def _check_sample_metadata(path: Path, sample_ids: list[str], *, cancelled=None) -> None:
        try:
            with path.open(encoding="utf-8", newline="") as handle:
                reader = csv.reader(handle, delimiter="\t")
                header = next(reader, None)
                if not header or header[0].strip() != "sample_id":
                    raise ProjectLoadError("项目样本信息表缺少 sample_id 列。")
                keys = []
                for line_number, row in enumerate(reader, start=2):
                    if line_number % 1024 == 0:
                        ProjectStore._raise_if_cancelled(cancelled)
                    if row:
                        keys.append(row[0])
        except (OSError, UnicodeError, csv.Error, StopIteration) as exc:
            raise ProjectLoadError("项目样本信息表损坏或无法读取。") from exc
        if keys != sample_ids:
            raise ProjectLoadError("样本信息编号或顺序与表达矩阵不一致。")

    @staticmethod
    def _raise_if_cancelled(cancelled) -> None:
        if cancelled is not None and cancelled():
            raise ProjectLoadCancelled("已取消项目校验；没有启动新分析。若检测到中断事务，软件已按事务记录恢复。")

    @classmethod
    def _sha256_file(cls, path: Path, *, progress=None, cancelled=None) -> str:
        digest = hashlib.sha256()
        total = path.stat().st_size
        completed = 0
        last_reported = 0
        if progress is not None:
            progress(0, total)
        with path.open("rb") as handle:
            while True:
                cls._raise_if_cancelled(cancelled)
                block = handle.read(1024 * 1024)
                if not block:
                    break
                digest.update(block)
                completed += len(block)
                if progress is not None and (
                    completed - last_reported >= 8 * 1024 * 1024 or completed == total
                ):
                    progress(completed, total)
                    last_reported = completed
        cls._raise_if_cancelled(cancelled)
        if progress is not None and completed != last_reported:
            progress(completed, total)
        return digest.hexdigest()

    @staticmethod
    def _column_count(path: Path) -> int:
        try:
            with path.open(encoding="utf-8", newline="") as handle:
                return len(next(csv.reader(handle, delimiter="\t")))
        except (OSError, UnicodeError, csv.Error, StopIteration) as exc:
            raise ProjectLoadError(f"项目表格无法读取：{path.name}") from exc

    @staticmethod
    def _validation_summary(
        raw: Any,
        feature_count: int,
        sample_count: int,
        numeric_count: int,
        warnings: list[str],
    ) -> ValidationSummary:
        if not isinstance(raw, dict):
            warnings.append("旧项目没有整理校验摘要；本次只完成结构和数值检查，不代表旧整理流程已通过验证。")
            return ValidationSummary(True, (), feature_count, sample_count, numeric_count)
        if raw.get("passed") is False:
            raise ProjectLoadError("项目整理校验明确标记为未通过；不能恢复该项目。")
        expected_counts = {
            "feature_count": feature_count,
            "sample_count": sample_count,
            "numeric_cell_count": numeric_count,
        }
        try:
            for key, actual in expected_counts.items():
                if key not in raw:
                    warnings.append(f"旧项目校验摘要缺少 {key}；已从矩阵结构重新统计。")
                elif int(raw[key]) != actual:
                    raise ProjectLoadError("项目清单中的数据量与矩阵文件不一致。")
        except (TypeError, ValueError, OverflowError) as exc:
            if isinstance(exc, ProjectLoadError):
                raise
            raise ProjectLoadError("项目清单中的数据量摘要无效。") from exc
        issues = []
        raw_issues = raw.get("issues", [])
        if not isinstance(raw_issues, list):
            raise ProjectLoadError("项目校验摘要中的问题列表格式无效。")
        for value in raw_issues:
            if not isinstance(value, dict):
                raise ProjectLoadError("项目校验摘要中的问题记录格式无效。")
            try:
                severity = IssueSeverity(str(value.get("severity", "INFO")))
            except ValueError as exc:
                raise ProjectLoadError("项目校验摘要包含未知严重级别。") from exc
            issues.append(Issue(
                str(value.get("code", "PROJECT_NOTE")), severity,
                str(value.get("message_zh", "")), str(value.get("operation_zh", "")),
            ))
        passed_recorded = raw.get("passed") is True
        if not passed_recorded:
            if any(issue.severity == IssueSeverity.BLOCKING for issue in issues):
                raise ProjectLoadError("旧项目校验摘要包含阻断问题；不能恢复该项目。")
            warnings.append("旧项目没有明确记录整理校验通过；当前仅基于文件结构恢复，且保持只读。")
        return ValidationSummary(True, tuple(issues), feature_count, sample_count, numeric_count)

    @classmethod
    def _load_plan(
        cls,
        root: Path,
        matrix_digest: str,
        feature_path: Path,
        sample_ids: list[str],
        warnings: list[str],
    ) -> dict[str, Any] | None:
        candidates = [root / "analysis" / "analysis-plan.json", root / "analysis-plan.json"]
        existing = [path for path in candidates if path.is_file()]
        if not existing:
            return None
        plans = []
        for path in existing:
            plan = cls._json(cls._member(root, str(path.relative_to(root))), path.name)
            if plan.get("matrix_hash") != matrix_digest:
                warnings.append(f"{path.name} 与当前矩阵哈希不匹配；未恢复该计划。")
                continue
            expected_metadata_hash = plan.get("feature_metadata_hash")
            if expected_metadata_hash and expected_metadata_hash != sha256_file(feature_path):
                warnings.append(f"{path.name} 的特征注释校验失败；未恢复该计划。")
                continue
            if not expected_metadata_hash:
                warnings.append(f"{path.name} 没有特征注释哈希；恢复后必须重新确认计划。")
            if not cls._valid_design(plan.get("design"), sample_ids):
                warnings.append(f"{path.name} 的分组或样本编号与当前矩阵不一致；未恢复该计划。")
                continue
            plans.append(plan)
        if len(plans) == 2 and plans[0] != plans[1]:
            warnings.append("项目中根目录与 analysis/ 下的计划内容不同；为避免选错，未自动恢复计划。")
            return None
        return plans[0] if plans else None

    @staticmethod
    def _valid_design(design: Any, sample_ids: list[str]) -> bool:
        if not isinstance(design, dict):
            return False
        assignments = design.get("assignments")
        if not isinstance(assignments, dict) or set(assignments) != set(sample_ids):
            return False
        if any(
            not isinstance(value, str) or value not in {"control", "test", "excluded"}
            for value in assignments.values()
        ):
            return False
        for key in ("paired_subjects", "batch"):
            mapping = design.get(key, {})
            if (
                not isinstance(mapping, dict)
                or not set(mapping).issubset(sample_ids)
                or any(not isinstance(value, str) for value in mapping.values())
            ):
                return False
        return True

    @classmethod
    def _load_result(
        cls,
        root: Path,
        feature_ids: list[str],
        matrix_digest: str,
        feature_path: Path,
        sample_path: Path,
        saved_plan: dict[str, Any] | None,
        warnings: list[str],
    ) -> tuple[dict[str, Any] | None, Path | None, bool, bool]:
        result_path = root / "analysis" / "deg-result.json"
        result_exists = result_path.is_file()
        result = cls._json(cls._member(root, "analysis/deg-result.json"), result_path.name) if result_exists else None
        report_candidate = root / "reports" / "analysis-report.html"
        report_path = cls._member(root, "reports/analysis-report.html") if report_candidate.is_file() else None
        manifest_path = root / "reports" / "analysis-manifest.json"
        report_valid = False
        result_integrity_valid = False
        result_hash_mismatch = False
        if manifest_path.is_file():
            report_manifest = cls._json(cls._member(root, "reports/analysis-manifest.json"), manifest_path.name)
            checks = {
                "matrix_hash": matrix_digest,
                "feature_metadata_hash": sha256_file(feature_path),
                "sample_metadata_hash": sha256_file(sample_path),
                "preparation_recipe_hash": sha256_file(root / "preparation-recipe.json")
                if (root / "preparation-recipe.json").is_file() else None,
                "analysis_plan_hash": sha256_file(root / "analysis" / "analysis-plan.json")
                if (root / "analysis" / "analysis-plan.json").is_file() else None,
                "deg_result_hash": sha256_file(result_path) if result_exists else None,
                "html_hash": sha256_file(report_path) if report_path is not None else None,
                "xlsx_hash": sha256_file(cls._member(root, "reports/deg-results.xlsx"))
                if (root / "reports" / "deg-results.xlsx").is_file() else None,
            }
            critical_result_keys = {
                "matrix_hash", "feature_metadata_hash", "sample_metadata_hash",
                "preparation_recipe_hash", "analysis_plan_hash", "deg_result_hash",
            }
            result_integrity_valid = True
            report_valid = report_path is not None
            for key, actual in checks.items():
                expected = report_manifest.get(key)
                if expected is None:
                    report_valid = False
                    if key in critical_result_keys:
                        result_integrity_valid = False
                    warnings.append(f"旧报告清单缺少 {key}；相关产物完整性无法验证。")
                elif actual is None or str(expected) != actual:
                    report_valid = False
                    if key in critical_result_keys:
                        result_hash_mismatch = True
                        result_integrity_valid = False
                    warnings.append(f"既有报告的 {key} 校验失败；相关旧产物不会作为匹配当前项目的结果打开。")
            expected_filter_hash = report_manifest.get("result_filter_hash")
            if expected_filter_hash is not None:
                filter_path = root / "analysis" / "result-filter.json"
                if (
                    not filter_path.is_file()
                    or filter_path.is_symlink()
                    or sha256_file(filter_path) != str(expected_filter_hash)
                ):
                    report_valid = False
                    result_integrity_valid = False
                    result_hash_mismatch = True
                    warnings.append("结果筛选设置与报告清单中的校验值不一致；报告不作为匹配当前项目的结果打开。")
            expected_annotation_hash = report_manifest.get("gene_annotations_hash")
            if expected_annotation_hash is not None:
                annotation_path = root / "analysis" / "gene-annotations.tsv"
                if (
                    annotation_path.is_symlink()
                    or not annotation_path.is_file()
                    or sha256_file(annotation_path) != str(expected_annotation_hash)
                ):
                    report_valid = False
                    result_integrity_valid = False
                    result_hash_mismatch = True
                    warnings.append("后期基因注释与报告清单中的校验值不一致；报告不会作为匹配当前项目的结果打开。")
            figure_manifest_path = root / "figures" / "figure-manifest.json"
            expected_figure_manifest_hash = report_manifest.get("figure_manifest_hash")
            if expected_figure_manifest_hash is None:
                report_valid = False
                warnings.append("报告清单没有图表清单校验值；图表与结果的一致性未验证。")
            elif (
                not figure_manifest_path.is_file()
                or figure_manifest_path.is_symlink()
                or sha256_file(figure_manifest_path) != str(expected_figure_manifest_hash)
            ):
                report_valid = False
                warnings.append("图表清单校验失败；不会把 HTML 报告作为完整一致的结果打开。")
            figure_issues = cls._validate_figures(root, report_manifest)
            if figure_issues:
                report_valid = False
                warnings.extend(figure_issues)
            if result is not None and saved_plan is None:
                report_valid = False
                result_integrity_valid = False
                warnings.append("报告有结果校验值，但对应分析计划未能恢复；为避免误配，不打开该报告。")
        else:
            warnings.append("没有分析报告清单；既有 HTML 报告无法与当前项目核对。")

        if result is not None:
            rows = result.get("table")
            required_row_fields = {
                "feature_key", "original_id", "gene_symbol", "log2FoldChange", "pvalue", "padj",
                "direction", "tested", "reason_not_tested", "mean_control_expression",
                "mean_treated_expression", "analysis_method", "analysis_basis", "effect_scale",
                "mean_expression_unit",
            }
            if (
                not isinstance(rows, list)
                or len(rows) != len(feature_ids)
                or any(not isinstance(row, dict) or not required_row_fields.issubset(row) for row in rows)
                or [str(row.get("feature_key", "")) for row in rows if isinstance(row, dict)] != feature_ids
            ):
                warnings.append("既有结果与矩阵的特征编号/顺序不一致；未加载结果。")
                result = None
            elif saved_plan is None:
                warnings.append("既有结果没有可恢复且匹配的分析计划；仅保留文件，不展示为已恢复结果。")
                result = None
            elif any(
                not isinstance(row.get("tested"), bool)
                or row.get("analysis_method") != saved_plan.get("method")
                or row.get("analysis_basis") != saved_plan.get("analysis_basis")
                for row in rows
            ):
                warnings.append("既有结果行的方法或分析基础与已保存计划不一致；未加载结果。")
                result = None
            elif (
                result.get("method", saved_plan.get("method")) != saved_plan.get("method")
                or result.get("profile_id", saved_plan.get("method_profile_id"))
                != saved_plan.get("method_profile_id")
            ):
                warnings.append("既有结果的整体方法标识与已保存计划不一致；未加载结果。")
                result = None
            elif result_hash_mismatch:
                warnings.append("既有结果文件与报告清单中的校验值不同；未加载被改动的结果。")
                result = None
            elif not result_integrity_valid:
                warnings.append("既有结果缺少完整校验值或与当前项目不匹配；仅保留为未完整核验的历史记录，不作为已验证结果使用。")
        if result is None and report_valid:
            report_valid = False
            warnings.append("既有结果未通过恢复校验；为避免展示不一致内容，离线 HTML 报告已禁用。")
        if report_path is not None and not cls._valid_html_report(report_path):
            report_valid = False
            warnings.append("离线 HTML 文件不是完整、非空的报告文档；仍保留既有结果供只读查看。")
        xlsx_path = root / "reports" / "deg-results.xlsx"
        if result is not None and not cls._valid_result_workbook(xlsx_path, result.get("table")):
            report_valid = False
            warnings.append("XLSX 工作簿内容无法打开，或与结果 JSON 的五列导出不一致；仍保留既有结果供只读查看。")
        return result, report_path, report_valid, result_integrity_valid

    @classmethod
    def _valid_result_workbook(cls, path: Path, rows: Any) -> bool:
        """检查唯一工作表及其固定五列是否逐行对应结果 JSON。"""
        if not path.is_file() or path.is_symlink() or not isinstance(rows, list):
            return False
        try:
            from openpyxl import load_workbook

            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                if workbook.sheetnames != ["DEG Results"]:
                    return False
                sheet = workbook["DEG Results"]
                if sheet.max_column != len(_RESULT_EXPORT_COLUMNS) or sheet.max_row != len(rows) + 1:
                    return False
                values = sheet.iter_rows(min_row=1, max_row=len(rows) + 1, values_only=True)
                if next(values, None) != _RESULT_EXPORT_COLUMNS:
                    return False
                for actual_row, result_row in zip(values, rows):
                    if not isinstance(result_row, dict) or len(actual_row) != len(_RESULT_EXPORT_COLUMNS):
                        return False
                    expected = tuple(cls._exported_value(result_row.get(column)) for column in _RESULT_EXPORT_COLUMNS)
                    actual = tuple(cls._exported_value(value) for value in actual_row)
                    if any(
                        not cls._same_exported_value(left, right)
                        for left, right in zip(actual, expected)
                    ):
                        return False
                return True
            finally:
                workbook.close()
        except Exception:
            # 产物损坏时只禁用离线报告，不影响已恢复结果的只读查看。
            return False

    @staticmethod
    def _exported_value(value: Any) -> Any:
        if value is None or (isinstance(value, str) and not value.strip()):
            return "NA"
        if isinstance(value, (int, float)) and not isinstance(value, bool) and not math.isfinite(value):
            return "NA"
        return value

    @staticmethod
    def _same_exported_value(actual: Any, expected: Any) -> bool:
        if (
            isinstance(actual, (int, float)) and not isinstance(actual, bool)
            and isinstance(expected, (int, float)) and not isinstance(expected, bool)
        ):
            # XLSX numeric cells have finite decimal precision; allow only the
            # serialization-sized rounding introduced by the spreadsheet format.
            return math.isclose(float(actual), float(expected), rel_tol=1e-14, abs_tol=1e-15)
        return type(actual) is type(expected) and actual == expected

    @staticmethod
    def _valid_html_report(path: Path) -> bool:
        if not path.is_file() or path.is_symlink():
            return False
        try:
            content = path.read_text(encoding="utf-8")
            if not content.strip():
                return False
            parser = _ReportHTMLValidator()
            parser.feed(content)
            parser.close()
            return parser.is_valid_report()
        except (OSError, UnicodeError, ValueError):
            return False

    @classmethod
    def _validate_figures(cls, root: Path, report_manifest: dict[str, Any]) -> list[str]:
        issues: list[str] = []
        try:
            figure_manifest = cls._json(
                cls._member(root, "figures/figure-manifest.json"), "图表清单 figure-manifest.json",
            )
        except (ProjectLoadError, OSError):
            return ["图表清单缺失或内容无效；不会把 HTML 报告作为完整结果打开。"]
        if set(figure_manifest) != set(_FIGURE_FILES):
            issues.append("图表清单没有恰好包含规定的五张图；不会把 HTML 报告作为完整结果打开。")
        for key, filename in _FIGURE_FILES.items():
            value = figure_manifest.get(key)
            if not isinstance(value, str) or Path(value).name != filename:
                issues.append(f"图表清单中的 {key} 与规定文件名 {filename} 不匹配。")

        figure_hashes = report_manifest.get("figure_hashes")
        expected_names = set(_FIGURE_FILES.values())
        if not isinstance(figure_hashes, dict) or set(figure_hashes) != expected_names:
            issues.append("报告清单没有恰好包含五张规定图表的校验值。")
            figure_hashes = figure_hashes if isinstance(figure_hashes, dict) else {}
        for filename in _FIGURE_FILES.values():
            expected_hash = figure_hashes.get(filename)
            try:
                figure_path = cls._member(root, f"figures/{filename}")
                hash_matches = isinstance(expected_hash, str) and sha256_file(figure_path) == expected_hash
                if not hash_matches:
                    issues.append(f"图表 {filename} 缺失或校验失败；不会把 HTML 报告作为完整结果打开。")
                    continue
                from PIL import Image

                with Image.open(figure_path) as image:
                    image.verify()
                if image.width <= 0 or image.height <= 0:
                    raise ValueError("图像尺寸无效")
            except Exception:
                issues.append(f"图表 {filename} 无法解码或校验失败；不会把 HTML 报告作为完整结果打开。")
        return issues

    @staticmethod
    def _measurement_provenance(confirmations: Any) -> str:
        if not isinstance(confirmations, list):
            return "user-selected input facts: no confirmed values"
        entries = []
        for item in confirmations:
            if not isinstance(item, dict):
                continue
            key = str(item.get("fact_key", "")).strip()
            value = str(item.get("selected_value", "")).strip()
            if not key or not value or value == "UNKNOWN":
                continue
            status = str(item.get("status", ""))
            source = {
                "user_confirmed_without_direct_text": "manual_without_direct_text",
                "user_confirmed_text_candidate": "linked_text_candidate",
                "user_resolved_conflicting_text": "linked_text_candidate",
                "user_overrode_text_candidate": "manual_override_of_text_candidate",
            }.get(status, status or "unspecified_confirmation")
            entries.append(f"{key}={value}[{source}]")
        return "user-selected input facts: " + ("; ".join(entries) if entries else "no confirmed values")
