from __future__ import annotations

import html
import json
import math
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from rnaseq_helper.domain.models import AnalysisPlan, ExpressionMatrix, sha256_file
from rnaseq_helper.analysis.result_filter import passes_result_filter, selected_pvalue, summarize_result
from rnaseq_helper.domain.models import ResultFilterConfig


class ReportService:
    """生成不依赖网络的 HTML 和 XLSX 结果包。"""

    RESULT_EXPORT_COLUMNS = (
        "feature_key",
        "gene_symbol",
        "log2FoldChange",
        "pvalue",
        "padj",
    )
    MISSING_VALUE = "NA"

    def build(
        self,
        matrix: ExpressionMatrix,
        plan: AnalysisPlan,
        result: dict[str, Any],
        figures: dict[str, Path] | None = None,
        output_dir: Path | None = None,
        project_dir: Path | None = None,
        analysis_artifact_dir: Path | None = None,
    ) -> dict[str, Path]:
        project_dir = Path(project_dir) if project_dir is not None else matrix.values.path.parent.parent
        report_dir = Path(output_dir) if output_dir is not None else project_dir / "reports"
        report_dir.mkdir(parents=True, exist_ok=True)
        html_path = report_dir / "analysis-report.html"
        xlsx_path = report_dir / "deg-results.xlsx"
        manifest_path = report_dir / "analysis-manifest.json"
        interpretation_path = report_dir / "interpretation-zh.txt"
        interpretation = self._interpretation(matrix, plan, result)
        project_manifest = self._read_json(project_dir / "project.json")
        analysis_artifact_dir = (
            Path(analysis_artifact_dir)
            if analysis_artifact_dir is not None
            else project_dir / "analysis"
        )
        self._write_html(html_path, matrix, plan, result, figures or {}, interpretation, project_manifest)
        self._write_xlsx(xlsx_path, matrix, plan, result, interpretation, project_manifest)
        interpretation_path.write_text(interpretation + "\n", encoding="utf-8")
        figure_hashes = {
            path.name: sha256_file(path)
            for path in (figures or {}).values()
        }
        figure_manifest_path = (
            next(iter((figures or {}).values())).parent / "figure-manifest.json"
            if figures else None
        )
        manifest = {
            "created_at": datetime.now(timezone.utc).isoformat(),
            "html": "reports/" + html_path.name,
            "xlsx": "reports/" + xlsx_path.name,
            "interpretation_zh": "reports/" + interpretation_path.name,
            "analysis_plan": "analysis/analysis-plan.json",
            "preparation_recipe": "preparation-recipe.json",
            "method_profile_id": plan.method_profile_id,
            "source_modality": matrix.source_modality or project_manifest.get("source_modality"),
            "matrix_hash": matrix.content_hash,
            "feature_metadata_hash": sha256_file(matrix.feature_metadata.path),
            "sample_metadata_hash": sha256_file(matrix.sample_metadata.path),
            "preparation_recipe_hash": (
                sha256_file(project_dir / "preparation-recipe.json")
                if (project_dir / "preparation-recipe.json").is_file() else None
            ),
            "analysis_plan_hash": sha256_file(analysis_artifact_dir / "analysis-plan.json"),
            "deg_result_hash": sha256_file(analysis_artifact_dir / "deg-result.json"),
            "result_filter_hash": (
                sha256_file(analysis_artifact_dir / "result-filter.json")
                if (analysis_artifact_dir / "result-filter.json").is_file() else None
            ),
            "gene_annotations_hash": (
                sha256_file(analysis_artifact_dir / "gene-annotations.tsv")
                if (analysis_artifact_dir / "gene-annotations.tsv").is_file() else None
            ),
            "html_hash": sha256_file(html_path),
            "xlsx_hash": sha256_file(xlsx_path),
            "figure_manifest_hash": (
                sha256_file(figure_manifest_path)
                if figure_manifest_path is not None and figure_manifest_path.is_file()
                else None
            ),
            "figure_hashes": figure_hashes,
            "source_path": project_manifest.get("source_path"),
            "platform_id": matrix.platform_id or project_manifest.get("platform_id"),
            "reference_result": project_manifest.get("reference_result"),
            "measurement_provenance": matrix.measurement_provenance,
            "figures": {
                key: "figures/" + path.name
                for key, path in (figures or {}).items()
            },
            "result_issues": result.get("issues", []),
        }
        manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        return {"html": html_path, "xlsx": xlsx_path, "interpretation_zh": interpretation_path, "manifest": manifest_path}

    @staticmethod
    def _write_html(path: Path, matrix: ExpressionMatrix, plan: AnalysisPlan, result: dict[str, Any], figures: dict[str, Path], interpretation: str, project_manifest: dict[str, Any]) -> None:
        active_filter = ResultFilterConfig.from_mapping(
            result.get("active_result_filter") or plan.result_filter.as_dict()
        )
        rows = [row for row in result.get("table", []) if isinstance(row, dict)]
        summary = result.get("display_summary") or summarize_result(rows, active_filter)
        summary_labels = {
            "feature_count": "结果特征总数", "tested_count": "完成统计检验的特征",
            "not_tested_count": "未进入统计检验的特征", "significant_count": "达到显著性阈值的特征",
            "significant_magnitude_count": "同时达到显著性与效应阈值的特征",
            "up_count": "实验组较高", "down_count": "对照组较高",
        }
        metric_specs = (
            ("tested_count", "完成统计检验"),
            ("significant_magnitude_count", "符合当前筛选"),
            ("up_count", "实验组较高"),
            ("down_count", "对照组较高"),
        )
        metric_cards = "".join(
            f'<div class="metric"><strong>{html.escape(str(summary.get(key, 0)))}</strong><span>{label}</span></div>'
            for key, label in metric_specs
        )
        summary_rows = "".join(
            f"<tr><th>{label}</th><td>{html.escape(str(summary.get(key, 0)))}</td></tr>"
            for key, label in summary_labels.items() if key in summary
        )
        limitations = "".join(f"<li>{html.escape(item)}</li>" for item in plan.limitations_zh)
        warnings = "".join(f"<li>{html.escape(item)}</li>" for item in plan.warnings) or "<li>无额外分析告警</li>"
        result_issues = "".join(
            f"<li>{html.escape(str(issue.get('message_zh', issue)))}</li>" if isinstance(issue, dict)
            else f"<li>{html.escape(str(issue))}</li>" for issue in result.get("issues", [])
        ) or "<li>未记录额外运行问题</li>"
        parameter_items = "".join(
            f"<li><code>{html.escape(str(key))}</code>：{html.escape(json.dumps(value, ensure_ascii=False, default=str) if isinstance(value, (dict, list, tuple)) else str(value))}</li>"
            for key, value in result.get("parameters", {}).items()
        ) or "<li>未记录</li>"
        diagnostic_items = "".join(
            f"<li>{html.escape(ReportService._diagnostic_label(key))}：{html.escape(ReportService._format_diagnostic_value(key, value))}</li>"
            for key, value in result.get("diagnostics", {}).items()
        ) or "<li>未记录额外诊断</li>"
        runtime_items = "".join(
            f"<li>{html.escape(str(key))}：{html.escape(str(value))}</li>"
            for key, value in result.get("runtime", {}).items()
        ) or "<li>未记录运行环境版本</li>"
        contrast = plan.contrast if isinstance(plan.contrast, dict) else {}
        control_group = str(contrast.get("control", "")).strip()
        test_group = str(contrast.get("test", "")).strip()
        contrast_description = (
            f"实验组 {test_group} − 对照组 {control_group}（log2FoldChange > 0 表示实验组在模型尺度上较高）"
            if control_group and test_group
            else "对比方向：分析计划未记录完整的测试组与对照组，无法从报告确认"
        )
        transform_description = ReportService._transform_description(plan, project_manifest)
        assignments = plan.design.get("assignments", {})
        subjects = plan.design.get("paired_subjects", {})
        batches = plan.design.get("batch", {})
        subject_aliases = ReportService._factor_aliases(subjects, "个体")
        batch_aliases = ReportService._factor_aliases(batches, "批次")
        assignment_items = ""
        for group_key, group_title in ((control_group, "对照组"), (test_group, "实验组")):
            if not group_key:
                continue
            members = [sample for sample, group in assignments.items() if group == group_key]
            member_items = "".join(
                "<li>" + html.escape(str(sample))
                + (f"；配对个体 {html.escape(subject_aliases[subjects[sample]])}"
                   if sample in subjects and str(subjects[sample]).strip() else "")
                + (f"；{html.escape(batch_aliases[batches[sample]])}"
                   if sample in batches and str(batches[sample]).strip() else "")
                + "</li>" for sample in members
            ) or "<li>未记录样本</li>"
            assignment_items += (
                f'<div class="card"><h3>{group_title}（{len(members)} 个）</h3>'
                f'<p class="muted">组别标记：{html.escape(str(group_key))}</p><ul>{member_items}</ul></div>'
            )
        excluded_samples = [
            sample for sample, group in assignments.items()
            if group not in {control_group, test_group}
        ]
        excluded_items = "".join(f"<li>{html.escape(str(sample))}</li>" for sample in excluded_samples)
        assignment_items += (
            f'<div class="card"><h3>未纳入分析（{len(excluded_samples)} 个）</h3>'
            f'<ul>{excluded_items or "<li>无</li>"}</ul></div>'
        )
        replicate_structure_labels = {
            "independent": "已确认不同个体/独立生物样本",
            "paired": "已确认同一对象的配对/重复测量",
            "cell": "细胞/细胞核级（当前版本不支持直接分析）",
            "unknown": "未知或未确认",
        }
        replicate_structure = replicate_structure_labels.get(
            str(plan.design.get("replicate_structure", "")), "旧计划未记录样本独立性/配对确认",
        )
        factor_privacy_note = (
            "<p>报告中的配对个体与批次已用匿名序号替代原标签；相同原标签仍对应相同序号。</p>"
            if subject_aliases or batch_aliases else ""
        )
        contrast_groups = tuple(
            group for group in (plan.contrast.get("control"), plan.contrast.get("test"))
            if group not in (None, "")
        )
        included_counts = {
            group: sum(assigned_group == group for assigned_group in assignments.values())
            for group in contrast_groups
        }
        excluded_count = len(excluded_samples)
        if assignments and contrast_groups:
            analyzed_samples = sum(included_counts.values())
            sample_usage = (
                f"分析纳入样本：{analyzed_samples}（"
                + "；".join(
                    f"{('对照组' if group == control_group else '实验组')} {count}"
                    for group, count in included_counts.items()
                )
                + f"）；未纳入 {excluded_count} 个样本。"
            )
        else:
            sample_usage = "分析纳入样本：无法从保存的分组计划确认。"
        figure_labels = {
            "pca": ("样本整体差异（PCA）", "主成分压缩展示变换后表达的主要变异；点彼此接近表示整体表达谱相似，不代表因果或显著性。"),
            "sample_correlation": ("样本相关性", "展示样本两两表达模式的相关程度；相关性高不等于生物学重复或实验设计正确。"),
            "pvalue_histogram": ("原始 P 值分布", "仅展示完成检验特征的原始 P 值分布；应结合设计与多重检验结果解读，不能单独作为质量结论。"),
            "volcano": ("差异效应与显著性（火山图）", "横轴为 log2FoldChange，纵轴为 −log10(P)；显著性使用当前筛选所选 P 值类型。"),
            "top_heatmap": ("优先差异特征表达热图", "展示按当前筛选选出的特征及模型使用的变换后表达；行内标准化用于模式比较，不表示绝对表达量。"),
        }
        figure_sections = []
        for key, figure_path in figures.items():
            title, caption = figure_labels.get(key, (key.replace("_", " "), "图形用于辅助探索，不替代统计检验与实验设计核对。"))
            rel = "../figures/" + html.escape(figure_path.name, quote=True)
            figure_sections.append(
                f'<figure><img src="{rel}" alt="{html.escape(title, quote=True)}">'
                f'<figcaption><strong>{html.escape(title)}</strong><p>{html.escape(caption)}</p>'
                f'<a href="{rel}">单独打开原图</a></figcaption></figure>'
            )
        figure_html = "".join(figure_sections) or '<p class="muted">本项目未提供图表文件。</p>'
        provenance_items = "".join(
            f"<li>{html.escape(str(key))}：{html.escape(str(value))}</li>"
            for key, value in {
                "来源文件名": Path(str(project_manifest.get("source_path", ""))).name,
                "平台": matrix.platform_id or project_manifest.get("platform_id", ""),
                "来源平台类别": matrix.source_modality or project_manifest.get("source_modality", ""),
                "参考结果": project_manifest.get("reference_result", "未使用本地参考表"),
                "单位/尺度选择记录": matrix.measurement_provenance or "未记录来源证据",
            }.items()
            if value
        )
        validation = project_manifest.get("validation", {})
        validation_items = "".join(
            f"<li>{html.escape(str(issue.get('severity', 'INFO')))}：{html.escape(str(issue.get('message_zh', issue)))}</li>"
            if isinstance(issue, dict) else f"<li>{html.escape(str(issue))}</li>"
            for issue in validation.get("issues", [])
        ) or "<li>整理验证无额外记录</li>"
        selected_rows = [row for row in rows if passes_result_filter(row, active_filter)]

        def representative_rows(direction: str) -> str:
            directional = [
                row for row in selected_rows
                if (float(row["log2FoldChange"]) > 0 if direction == "positive" else float(row["log2FoldChange"]) < 0)
            ]
            directional.sort(key=lambda row: (
                selected_pvalue(row, active_filter) is None,
                selected_pvalue(row, active_filter) or 0.0,
                str(row.get("feature_key", "")),
            ))
            top_rows = directional[:10]
            return "".join(
                "<tr>" + "".join(
                    f"<td>{ReportService._html_result_value(row.get(key))}</td>"
                    for key in ReportService.RESULT_EXPORT_COLUMNS
                ) + "</tr>" for row in top_rows
            ) or '<tr><td colspan="5">当前筛选条件下没有符合方向的特征。</td></tr>'

        positive_rows_html = representative_rows("positive")
        negative_rows_html = representative_rows("negative")
        p_label = "FDR（BH 校正后 P 值）" if active_filter.pvalue_mode == "fdr" else "原始 P 值（未校正）"
        selected_clause = f"{p_label} &lt; {active_filter.pvalue_threshold:g}，且 |log2FoldChange| ≥ {active_filter.log2fc_threshold:g}"
        analysis_name = str(project_manifest.get("analysis_name") or "差异表达分析")
        document = f"""<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><meta name="color-scheme" content="light"><title>{html.escape(analysis_name)}｜RNA-Seq Helper 分析报告</title>
<style>
:root{{color-scheme:light;--ink:#17212b;--muted:#536577;--line:#d7e0e7;--paper:#fff;--wash:#f3f7fa;--accent:#1668a8;--soft:#eaf3fa}}
*{{box-sizing:border-box}}html{{background:#edf1f4}}body{{font-family:-apple-system,BlinkMacSystemFont,"Helvetica Neue",Arial,sans-serif;color:var(--ink);background:var(--paper);max-width:1040px;margin:28px auto;padding:48px 62px;line-height:1.65;box-shadow:0 8px 32px #243b531c}}
h1{{font-size:30px;line-height:1.25;margin:0 0 10px}}h2{{font-size:20px;margin:34px 0 12px;padding-bottom:8px;border-bottom:1px solid var(--line)}}h3{{font-size:16px;margin:20px 0 8px}}p{{margin:8px 0}}.eyebrow{{color:var(--accent);font-weight:700;letter-spacing:.06em}}.lead{{font-size:16px;color:var(--muted)}}.meta{{color:var(--muted);font-size:14px}}.metrics{{display:grid;grid-template-columns:repeat(4,1fr);gap:12px;margin:18px 0}}.metric{{background:var(--wash);border:1px solid var(--line);border-radius:10px;padding:15px}}.metric strong{{display:block;font-size:26px;color:var(--accent)}}.metric span{{font-size:13px;color:var(--muted)}}.callout{{background:var(--soft);border-left:4px solid var(--accent);padding:14px 16px;margin:18px 0;border-radius:4px}}.grid{{display:grid;grid-template-columns:1fr 1fr;gap:14px}}.card{{border:1px solid var(--line);border-radius:10px;padding:16px;background:#fff}}table{{width:100%;border-collapse:collapse;margin:12px 0;font-size:14px}}th,td{{text-align:left;vertical-align:top;padding:9px 11px;border-bottom:1px solid var(--line)}}th{{background:var(--wash)}}.result-table{{font-size:12px;table-layout:fixed;overflow-wrap:anywhere}}.result-table th:nth-child(1){{width:24%}}.result-table th:nth-child(2){{width:20%}}.result-table th:nth-child(n+3){{width:18.6%}}figure{{margin:18px 0 26px;break-inside:avoid;border:1px solid var(--line);border-radius:10px;padding:14px}}figure img{{display:block;max-width:100%;height:auto;margin:auto}}figcaption{{margin-top:12px;color:var(--muted);font-size:14px}}figcaption strong{{color:var(--ink);font-size:16px}}figcaption p{{margin:4px 0}}a{{color:var(--accent)}}.muted{{color:var(--muted)}}pre{{white-space:pre-wrap;background:var(--wash);padding:16px;border-radius:8px;font:inherit}}details{{border:1px solid var(--line);border-radius:10px;padding:14px 16px;margin-top:20px}}summary{{cursor:pointer;font-weight:700}}.technical{{font-size:13px;overflow-wrap:anywhere}}footer{{margin-top:40px;padding-top:14px;border-top:1px solid var(--line);font-size:12px;color:var(--muted)}}
@media(max-width:720px){{body{{margin:0;padding:26px 20px}}.metrics{{grid-template-columns:repeat(2,1fr)}}.grid{{grid-template-columns:1fr}}}}
@page{{size:A4;margin:15mm}}@media print{{html{{background:#fff}}body{{max-width:none;margin:0;padding:0;box-shadow:none;font-size:10.5pt}}h1{{font-size:24pt}}h2{{font-size:15pt;break-after:avoid}}h3{{break-after:avoid}}.card,.metric,figure,table{{break-inside:avoid}}a{{color:inherit;text-decoration:none}}.no-print{{display:none}}}}
</style></head><body>
<header><p class="eyebrow">RNA-Seq Helper · 离线分析报告</p><h1>{html.escape(analysis_name)}</h1><p class="lead">{html.escape(contrast_description)}</p><p class="meta">分析方法：{html.escape(plan.method)} · 样本关系：{html.escape(replicate_structure)}</p></header>
<section><h2>核心结果</h2><div class="metrics">{metric_cards}</div><p>当前筛选：{selected_clause}。</p><p>矩阵包含 {matrix.values.shape[0]} 个特征、{matrix.values.shape[1]} 列样本；{sample_usage}</p><p class="meta">统计结果保存在同目录的 <a href="deg-results.xlsx">deg-results.xlsx</a>（完整结果表，仅含约定的五列）。</p><div class="callout">模型效应是当前模型尺度上的组间差异，不等同于原始线性表达均值比，也不证明因果关系。未检测特征及缺失统计值不会被改写或删除。</div><table><tbody>{summary_rows}</tbody></table></section>
<section><h2>代表性结果</h2><p>按当前筛选条件分别展示 log2FoldChange 正值和负值的特征；每个方向按所选 P 值从小到大最多列出 10 个。完整统计结果请打开 XLSX。</p><div class="grid"><div class="card"><h3>log2FoldChange &gt; 0 · 实验组较高</h3><table class="result-table"><thead><tr>{''.join(f'<th>{html.escape(key)}</th>' for key in ReportService.RESULT_EXPORT_COLUMNS)}</tr></thead><tbody>{positive_rows_html}</tbody></table></div><div class="card"><h3>log2FoldChange &lt; 0 · 对照组较高</h3><table class="result-table"><thead><tr>{''.join(f'<th>{html.escape(key)}</th>' for key in ReportService.RESULT_EXPORT_COLUMNS)}</tr></thead><tbody>{negative_rows_html}</tbody></table></div></div></section>
<section><h2>图表与解读</h2>{figure_html}</section>
<section><h2>中文解读</h2><pre>{html.escape(interpretation)}</pre></section>
<section><h2>样本分组</h2><p>{html.escape(replicate_structure)}；仅对照组和实验组进入本次比较，其余样本不参与统计。</p><div class="grid">{assignment_items}</div>{factor_privacy_note}</section>
<section><h2>方法与限制</h2><p>{html.escape(plan.method_reason_zh)}</p><p>{html.escape(transform_description)}</p><ul>{limitations}</ul><p>完成统计检验：{summary.get('tested_count', 0)}；未进入检验：{summary.get('not_tested_count', 0)}。统计筛选只影响展示，不改变原始统计模型或统计列。</p></section>
<details><summary>技术附录：来源、参数、软件版本与质量记录</summary><div class="technical"><h3>来源与整理验证</h3><ul>{provenance_items or '<li>未记录额外来源信息</li>'}</ul><ul>{validation_items}</ul><h3>统计参数与运行环境</h3><ul>{parameter_items}</ul><ul>{runtime_items}</ul><h3>质量诊断</h3><ul>{diagnostic_items}</ul><ul>{warnings}</ul><ul>{result_issues}</ul><p>方法配置：<code>{html.escape(plan.method_profile_id)}</code></p><p>矩阵 SHA-256：<code>{html.escape(matrix.content_hash)}</code></p></div></details>
<footer>本报告为本地生成的静态 HTML，不加载网络资源。图表文件与本报告保持相对目录结构时可离线查看及打印。</footer></body></html>"""
        path.write_text(document, encoding="utf-8")

    @staticmethod
    def _write_xlsx(path: Path, matrix: ExpressionMatrix, plan: AnalysisPlan, result: dict[str, Any], interpretation: str, project_manifest: dict[str, Any]) -> None:
        workbook = Workbook()
        result_sheet = workbook.active
        result_sheet.title = "DEG Results"
        rows = result.get("table", [])
        result_sheet.append(list(ReportService.RESULT_EXPORT_COLUMNS))
        for row in rows:
            result_sheet.append([
                ReportService._export_result_value(row.get(column))
                for column in ReportService.RESULT_EXPORT_COLUMNS
            ])
        result_sheet.freeze_panes = "A2"
        result_sheet.auto_filter.ref = result_sheet.dimensions

        result_sheet.sheet_view.showGridLines = False
        header_fill = PatternFill("solid", fgColor="243B53")
        for cell in result_sheet[1]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for column_cells in result_sheet.columns:
            width = min(48, max(12, max(len(str(cell.value or "")) for cell in column_cells) + 2))
            result_sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width
        workbook.save(path)

    @classmethod
    def _export_result_value(cls, value: Any) -> Any:
        """将结果表中的缺失值写成可读的 NA，而不是空单元格。"""
        if value is None:
            return cls.MISSING_VALUE
        if isinstance(value, (int, float)) and not isinstance(value, bool) and not math.isfinite(value):
            return cls.MISSING_VALUE
        if isinstance(value, str) and not value.strip():
            return cls.MISSING_VALUE
        return value

    @classmethod
    def _html_result_value(cls, value: Any) -> str:
        """HTML 与 XLSX 对缺失值使用同一呈现规则，并进行 HTML 转义。"""
        return html.escape(str(cls._export_result_value(value)))

    @staticmethod
    def _read_json(path: Path) -> dict[str, Any]:
        try:
            value = json.loads(path.read_text(encoding="utf-8"))
            return value if isinstance(value, dict) else {}
        except (OSError, json.JSONDecodeError):
            return {}

    @staticmethod
    def _diagnostic_label(key: str) -> str:
        return {
            "residual_df": "残差自由度（每个特征）",
        }.get(key, key)

    @staticmethod
    def _factor_aliases(values: Any, prefix: str) -> dict[str, str]:
        aliases: dict[str, str] = {}
        for value in values.values() if isinstance(values, dict) else ():
            key = str(value).strip()
            if key and key not in aliases:
                aliases[key] = f"{prefix}{len(aliases) + 1}"
        return aliases

    @staticmethod
    def _format_diagnostic_value(key: str, value: Any) -> str:
        if key == "residual_df" and isinstance(value, (list, tuple)) and value:
            unique = list(dict.fromkeys(value))
            if len(unique) == 1:
                return f"{unique[0]}（所有特征相同）"
        if isinstance(value, (dict, list, tuple)):
            return json.dumps(value, ensure_ascii=False, default=str)
        return str(value)

    @staticmethod
    def _interpretation(matrix: ExpressionMatrix, plan: AnalysisPlan, result: dict[str, Any]) -> str:
        active_filter = result.get("active_result_filter")
        from rnaseq_helper.domain.models import ResultFilterConfig
        result_filter = ResultFilterConfig.from_mapping(active_filter or plan.result_filter.as_dict())
        summary = result.get("display_summary") or summarize_result(
            result.get("table", []), result_filter,
        )
        significant = int(summary.get("significant_magnitude_count", 0) or 0)
        tested = int(summary.get("tested_count", 0) or 0)
        direction = f"上调 {summary.get('up_count', 0)} 个、下调 {summary.get('down_count', 0)} 个"
        p_name = "FDR（BH 校正后 P 值）" if result_filter.pvalue_mode == "fdr" else "原始 P 值（未校正）"
        criterion = f"{p_name} < {result_filter.pvalue_threshold:g} 且 |模型效应|≥{result_filter.log2fc_threshold:g}"
        if significant:
            finding = f"按当前结果筛选规则 {criterion}，共有 {significant} 个特征达到条件（{direction}）。"
        else:
            finding = f"按当前结果筛选规则 {criterion}，没有特征同时达到条件。"
        basis = "原始 count 计数模型（DESeq2）" if plan.analysis_basis == "raw_count" else "标准化表达值的 log 表达模型（limma-trend）"
        contrast = plan.contrast if isinstance(plan.contrast, dict) else {}
        control_group = str(contrast.get("control", "")).strip()
        test_group = str(contrast.get("test", "")).strip()
        contrast_direction = (
            f"对比方向：{test_group} − {control_group}；正值表示测试组在模型尺度上高于对照组。"
            if control_group and test_group
            else "对比方向：分析计划未记录完整的测试组与对照组，方向无法核实。"
        )
        limitation = "本分析基于标准化值，未使用原始计数中的测序深度和计数精度信息。" if plan.analysis_basis != "raw_count" else "结果依赖原始 count 的实验设计、独立重复和异常值处理；NA 保留其未检验原因。"
        worker_warnings = [
            str(issue.get("message_zh", issue)) if isinstance(issue, dict) else str(issue)
            for issue in result.get("issues", [])
        ]
        warning_line = f"运行告警：{'；'.join(worker_warnings)}" if worker_warnings else "运行告警：无"
        return "\n".join([
            "本段为软件生成的可审计中文解读模板，不替代实验或临床判断。",
            f"分析基础：{basis}；本次有效检验特征数：{tested}。",
            contrast_direction,
            finding,
            "上述阈值只用于结果筛选，不会改变统计模型、原始 P 值、FDR 或模型效应；原始统计列均保留。",
            *( ["注意：使用未校正的原始 P 值筛选，未控制多重检验带来的假阳性风险。"] if result_filter.pvalue_mode == "pvalue" else [] ),
            "模型效应表示对比方向上的模型尺度差异，不直接等同于原尺度算术均值之比，也不证明因果关系。",
            limitation,
            warning_line,
            "建议结合 PCA、样本相关性、热图、实验设计和独立验证共同解释；软件不会因图形分离不明显而自动删除样本。",
        ])

    @staticmethod
    def _transform_description(plan: AnalysisPlan, project_manifest: dict[str, Any]) -> str:
        plan_preprocessing = getattr(plan, "preprocessing", {})
        preprocessing = plan_preprocessing if isinstance(plan_preprocessing, dict) else {}
        transform = str(preprocessing.get("transform", "")).strip()
        if transform == "log2(x+1)":
            threshold = preprocessing.get("filter_threshold")
            if threshold is None:
                return "软件本次对线性输入执行 log2(x+1) 变换；计划未记录过滤阈值。"
            return f"软件本次对线性输入先按 x≥{threshold} 过滤，再执行 log2(x+1)；未做额外重新标准化。"
        if transform == "already_log2":
            reported = str(project_manifest.get("reported_transform", "")).strip()
            if reported:
                return f"软件本次未再次取对数，直接拟合用户确认的 log2 尺度输入；来源文字记录的原始公式为：{reported}。"
            return "软件本次未再次取对数，直接拟合用户确认的 log2 尺度输入；项目未记录可核实的原文件公式，不能据此判断是否含伪计数。"
        if transform in {"none", ""} and plan.method == "DESEQ2":
            return "软件未对原始 count 做 log 变换；DESeq2 在计数模型中使用 size factor 处理样本测序深度差异。"
        if transform in {"none", ""}:
            return "软件本次未追加表达值变换；来源文件此前采用的变换公式未在此处推断。"
        return f"分析计划记录的变换标记为 {transform}；该标记不证明来源文件此前使用的具体公式。"
