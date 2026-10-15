from __future__ import annotations

import csv
import gzip
import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path
from typing import TYPE_CHECKING, Iterable

from openpyxl import load_workbook
import xlrd

from rnaseq_helper.domain.models import ReferenceRequest, ReferenceResult, ReferenceStatus
from rnaseq_helper.domain.models import JobContext
from rnaseq_helper.readers.gpl_reader import GPLAnnotationReader
from rnaseq_helper.readers.table_reader import TableReader

if TYPE_CHECKING:
    from rnaseq_helper.r_bridge.client import RBridgeClient


class ReferenceService:
    """优先使用用户提供的本地参考表，失败时保留原编号继续交付。"""

    CACHE_DIR = Path.home() / "Library" / "Application Support" / "RNA-Seq Helper" / "reference-cache"

    def __init__(self, bridge: "RBridgeClient | None" = None) -> None:
        self.bridge = bridge

    def resolve(self, request: ReferenceRequest) -> ReferenceResult:
        feature_ids = tuple(dict.fromkeys(str(value) for value in request.feature_ids if str(value)))
        if not feature_ids:
            return ReferenceResult(
                ReferenceStatus.NOT_NEEDED, {}, None, request.species_taxid, request.platform_id,
                None, 1.0, 0, 0, "没有需要映射的特征编号。", request=request,
                retrieved_at=datetime.now(timezone.utc).isoformat(),
            )
        source = next((path.expanduser().resolve() for path in request.source_paths if path.expanduser().is_file()), None)
        if source is None:
            orgdb_result = self._resolve_orgdb(request, feature_ids)
            if orgdb_result is not None:
                return orgdb_result
            return self._unavailable(
                request,
                feature_ids,
                "未提供可读取的本地参考表，且当前物种/编号没有可用的离线 OrgDb；本次仍保留原编号输出，不阻断按原编号分析。",
            )
        try:
            mapping = self._read_mapping(source)
        except (OSError, UnicodeError, ValueError, KeyError) as exc:
            return ReferenceResult(
                ReferenceStatus.UNAVAILABLE,
                {},
                "local_user_reference",
                request.species_taxid,
                request.platform_id,
                self._sha256(source),
                0.0,
                0,
                len(feature_ids),
                f"本地参考表读取失败（{exc}）；本次仍保留原编号输出，不阻断按原编号分析。",
                request=request,
                release="local-file",
                retrieved_at=datetime.now(timezone.utc).isoformat(),
            )
        selected = {feature_id: mapping[feature_id] for feature_id in feature_ids if feature_id in mapping}
        unmapped = len(feature_ids) - len(selected)
        ambiguous = sum(len(symbols) > 1 for symbols in selected.values())
        coverage = len(selected) / len(feature_ids)
        status = ReferenceStatus.READY if unmapped == 0 else ReferenceStatus.PARTIAL
        checksum = self._sha256(source)
        try:
            self._cache(source, request, checksum, selected)
        except OSError:
            # 缓存目录权限不应阻断本次本地分析；结果仍包含 source checksum。
            pass
        return ReferenceResult(
            status,
            selected,
            "local_user_reference",
            request.species_taxid,
            request.platform_id,
            checksum,
            coverage,
            ambiguous,
            unmapped,
            f"本地参考表覆盖 {coverage:.1%}；一对多映射保留为列表，不复制统计行。",
            request=request,
            release="local-file",
            retrieved_at=datetime.now(timezone.utc).isoformat(),
        )

    def _resolve_orgdb(self, request: ReferenceRequest, feature_ids: tuple[str, ...]) -> ReferenceResult | None:
        """尝试使用已安装的离线 OrgDb；任何不可用均返回 None 以便安全降级。"""

        feature_id_type = self._feature_id_type(feature_ids)
        if feature_id_type not in {"ENSEMBL_GENE", "ENTREZ_OR_NUMERIC", "SYMBOL"} or not request.species_taxid:
            return None
        try:
            from rnaseq_helper.r_bridge.client import RBridgeClient, RBridgeError

            bridge = self.bridge or RBridgeClient()
            response = bridge.run(
                "reference_orgdb",
                {
                    "feature_ids": list(feature_ids),
                    "species_taxid": request.species_taxid,
                    "id_type": feature_id_type,
                },
                JobContext(Path.cwd()),
                "offline-orgdb-reference-v1",
            )
        except (RBridgeError, OSError, ValueError):
            return None
        payload = dict(response.result or {})
        raw_mapping = payload.get("mapping")
        if not isinstance(raw_mapping, dict):
            return None
        mapping: dict[str, tuple[str, ...]] = {}
        for feature_id, symbols in raw_mapping.items():
            if isinstance(symbols, (list, tuple)):
                values = tuple(dict.fromkeys(str(symbol).strip() for symbol in symbols if str(symbol).strip()))
            else:
                values = (str(symbols).strip(),) if str(symbols).strip() else ()
            if values:
                mapping[str(feature_id)] = values
        if not mapping:
            return None
        selected = {feature_id: mapping[feature_id] for feature_id in feature_ids if feature_id in mapping}
        coverage = len(selected) / len(feature_ids)
        return ReferenceResult(
            ReferenceStatus.READY if len(selected) == len(feature_ids) else ReferenceStatus.PARTIAL,
            selected,
            "offline_orgdb",
            request.species_taxid,
            request.platform_id,
            None,
            coverage,
            sum(len(symbols) > 1 for symbols in selected.values()),
            len(feature_ids) - len(selected),
            f"使用本机已安装的离线 OrgDb（{payload.get('package', '未标明包名')} {payload.get('release', '')}）；一对多映射保留为列表，不复制统计行。",
            request=request,
            release=str(payload.get("release", "")) or None,
            retrieved_at=datetime.now(timezone.utc).isoformat(),
        )

    @staticmethod
    def _unavailable(request: ReferenceRequest, feature_ids: tuple[str, ...], guide: str) -> ReferenceResult:
        return ReferenceResult(
            ReferenceStatus.UNAVAILABLE,
            {},
            None,
            request.species_taxid,
            request.platform_id,
            None,
            0.0,
            0,
            len(feature_ids),
            guide,
            request=request,
            retrieved_at=datetime.now(timezone.utc).isoformat(),
        )

    @staticmethod
    def _feature_id_type(feature_ids: tuple[str, ...]) -> str:
        if not feature_ids:
            return "CUSTOM_ORIGINAL"
        ensembl = sum(value.upper().split(".", 1)[0].startswith(("ENSG", "ENSMUSG", "ENSRNOG")) for value in feature_ids)
        numeric = sum(value.isdigit() for value in feature_ids)
        symbols = sum(value.replace("_", "").isalnum() and not value.isdigit() and not value.upper().startswith("ENS") for value in feature_ids)
        total = len(feature_ids)
        if ensembl / total >= 0.8:
            return "ENSEMBL_GENE"
        if numeric / total >= 0.8:
            return "ENTREZ_OR_NUMERIC"
        if symbols / total >= 0.8:
            return "SYMBOL"
        return "CUSTOM_ORIGINAL"

    @classmethod
    def _read_mapping(cls, path: Path) -> dict[str, tuple[str, ...]]:
        if path.suffix.lower() in {".soft", ".annot", ".bgx"} or (
            path.suffix.lower() == ".gz" and GPLAnnotationReader.looks_like(path)
        ):
            return GPLAnnotationReader().read(path).mapping
        rows: Iterable[tuple[object, ...]]
        if path.suffix.lower() in {".xlsx", ".xlsm"}:
            workbook = load_workbook(path, read_only=True, data_only=True)
            sheet = workbook.active
            rows = sheet.iter_rows(values_only=True)
        elif path.suffix.lower() == ".xls":
            # 供应商常把制表符文本伪装成 .xls；统一交给 TableReader，
            # 它会先识别这类文本，真正的 BIFF .xls 仍由 xlrd 读取。
            return cls._mapping_from_rows(TableReader().read(path).rows)
        else:
            delimiter = "," if path.suffix.lower() == ".csv" else "\t"
            if path.suffix.lower() == ".gz":
                with gzip.open(path, "rt", encoding="utf-8-sig", newline="") as handle:
                    return cls._mapping_from_rows(csv.reader(handle, delimiter=delimiter))
            with path.open(encoding="utf-8-sig", newline="") as handle:
                return cls._mapping_from_rows(csv.reader(handle, delimiter=delimiter))
        return cls._mapping_from_rows(rows)

    @classmethod
    def _mapping_from_rows(cls, rows: Iterable[Iterable[object]]) -> dict[str, tuple[str, ...]]:
        iterator = iter(rows)
        try:
            header = [str(value or "").strip().casefold() for value in next(iterator)]
        except StopIteration:
            return {}
        id_index = cls._find_column(
            header,
            {"feature_key", "original_id", "probe_id", "id", "ensembl", "entrez", "gene_id", "gene id"},
        )
        symbol_index = cls._find_column(
            header,
            {"gene_symbol", "symbol", "gene", "gene_name", "gene name", "gene symbol"},
        )
        if id_index is None or symbol_index is None:
            return {}
        result: dict[str, list[str]] = {}
        for row in iterator:
            values = list(row)
            if max(id_index, symbol_index) >= len(values):
                continue
            feature_id = str(values[id_index] or "").strip()
            symbol = str(values[symbol_index] or "").strip()
            if feature_id and symbol and symbol not in result.setdefault(feature_id, []):
                result[feature_id].append(symbol)
        return {feature_id: tuple(symbols) for feature_id, symbols in result.items()}

    @staticmethod
    def _find_column(header: list[str], names: set[str]) -> int | None:
        for index, value in enumerate(header):
            if value in names:
                return index
        return None

    @staticmethod
    def _sha256(path: Path) -> str:
        digest = hashlib.sha256()
        with path.open("rb") as handle:
            for block in iter(lambda: handle.read(1024 * 1024), b""):
                digest.update(block)
        return digest.hexdigest()

    @classmethod
    def _cache(cls, source: Path, request: ReferenceRequest, checksum: str, mapping: dict[str, tuple[str, ...]]) -> None:
        cls.CACHE_DIR.mkdir(parents=True, exist_ok=True)
        key = hashlib.sha256(f"{checksum}|{request.species_taxid}|{request.platform_id}".encode()).hexdigest()[:24]
        path = cls.CACHE_DIR / f"{key}.json"
        path.write_text(
            json.dumps({"source": str(source), "checksum": checksum, "species_taxid": request.species_taxid, "platform_id": request.platform_id, "mapping": mapping}, ensure_ascii=False, indent=2, default=list) + "\n",
            encoding="utf-8",
        )
